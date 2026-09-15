from datetime import date
from django.db.models import Q, Exists, OuterRef
from calendar import monthrange
from django.db import DatabaseError, IntegrityError
from api.models import SupervsiorLocation
import uuid
from api.models import SupervisorLogin
from rest_framework.decorators import api_view
from .services.uptime_service import get_lambda_uptime, get_rds_uptime,get_lambda_uptime_by_range
from datetime import date, timedelta
from django.shortcuts import render
from rest_framework.response import Response
import requests
<<<<<<< HEAD
=======
# Import Python's built-in date utilities.
#
# date:
# Used to get today's date and work with calendar dates.
#
# timedelta:
# Used to add or subtract a specific number of days from a date.
from datetime import date, timedelta


# Import Django's JsonResponse class.
#
# JsonResponse allows this API endpoint to return Python data
# as a JSON response that can be consumed by the React frontend.
from django.http import JsonResponse


# Import the Consumers model.
#
# This model is connected to the `readingmaster` database table
# and is used to retrieve consumer and meter reading data.
from api.models import Consumers

>>>>>>> b85f512 (Connect dashboard reading cycle to backend API)
from rest_framework.decorators import api_view, permission_classes
from .models import Consumers, MeterReaderRegistration, Office, SupervisorLogin, UserManagement
from .serializers import (
    MeterReaderRegistrationSerializer,
    ConsumerDataSerializer,
    ConsumerWiseDetailsSerializer,
    MridSerializer,
    Serail,
    ConsumersMeterRegistration,
    SupervisorLoginSerializer,
    UserManagementSerializer,
)
from django.db.models import Q
from rest_framework import status
import datetime
from .serializers import ConsumerSerializer
from rest_framework.decorators import parser_classes

from rest_framework.parsers import MultiPartParser, FormParser
from django.db import connection
import json
import jwt
import requests
from django.contrib.auth import authenticate
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.pagination import LimitOffsetPagination

from .serializers import FailedImageSerializer
from django.db import transaction
from django.utils.dateparse import parse_date

from django.http import JsonResponse
from django.db.models import Count, Case, When, IntegerField, F

# from decouple import config
from django.db.models import Q
from rest_framework.pagination import PageNumberPagination
from django.core.paginator import Paginator
import base64
from datetime import datetime, timedelta, date
from copy import deepcopy
import math
from rest_framework.views import APIView
from rest_framework.response import Response
from openpyxl import Workbook, load_workbook
from django.http import HttpResponse
from geojson import Point, Feature, FeatureCollection
from pyproj import CRS

from django.forms.models import model_to_dict


# import xlswriter
from django_filters import FilterSet

# Create your views here.

SECRETKEY = "6AZJYQ2T317WGPXC0UHVLDOR49FIBS8N5ME"

<<<<<<< HEAD
@parser_classes([MultiPartParser, FormParser])
@api_view(["POST"])
=======
# ============================================================
# DASHBOARD — READING CYCLE API
# ============================================================

def dashboard_reading_cycle(request):
    """
    Return Reading Cycle metrics for the dashboard.

    The current implementation treats one calendar month as
    one reading cycle.

    Returned values:

    - total:
      Number of consumers with completed readings.

    - expected:
      Total number of consumers expected in the current cycle.

    - dayOfCycle:
      Current day of the reading cycle.

    - cycleDays:
      Total number of days in the current cycle.

    - workingLeft:
      Number of remaining days in the cycle.

    - actualPace:
      Average number of readings completed per day during
      the most recent 7 days.
    """

    # --------------------------------------------------
    # GET TODAY'S DATE
    # --------------------------------------------------

    # Get the current server date.
    #
    # Example:
    # 11 September 2026
    today = date.today()


    # ==================================================
    # DETERMINE THE CURRENT READING CYCLE
    # ==================================================

    # The reading cycle currently follows the calendar month.
    #
    # Example:
    # If today is 11 Sep 2026:
    #
    # cycle_start → 01 Sep 2026
    cycle_start = today.replace(day=1)


    # --------------------------------------------------
    # CALCULATE THE FIRST DAY OF THE NEXT MONTH
    # --------------------------------------------------

    # December requires special handling because the next
    # month belongs to the following year.
    if today.month == 12:

        # Example:
        #
        # 01 Dec 2026
        #      ↓
        # 01 Jan 2027
        next_month = today.replace(
            year=today.year + 1,
            month=1,
            day=1
        )

    else:

        # For every other month, move to the first day
        # of the next month.
        #
        # Example:
        #
        # September
        #     ↓
        # 01 October
        next_month = today.replace(
            month=today.month + 1,
            day=1
        )


    # --------------------------------------------------
    # CALCULATE THE LAST DAY OF THE CURRENT CYCLE
    # --------------------------------------------------

    # The last day of the current month is one day before
    # the first day of the next month.
    #
    # Example:
    #
    # next_month → 01 Oct 2026
    #
    # cycle_end → 30 Sep 2026
    cycle_end = next_month - timedelta(days=1)


    # ==================================================
    # EXPECTED CONSUMERS
    # ==================================================

    # Retrieve all consumer records belonging to the
    # current billing month.
    #
    # bill_month_dt is a DateField in the Consumers model,
    # making it suitable for date-based filtering.
    #
    # The resulting count represents the total number of
    # consumers expected to be processed during this cycle.
    expected = Consumers.objects.filter(

        # Include records from the first day of the cycle.
        bill_month_dt__gte=cycle_start,

        # Include records up to the final day of the cycle.
        bill_month_dt__lte=cycle_end,

    ).count()


    # ==================================================
    # COMPLETED READINGS
    # ==================================================

    # Count consumer records that belong to the current
    # billing cycle and have a completed reading date.
    #
    # A non-null reading_date_db is currently treated as
    # an indication that a meter reading was completed.
    total = Consumers.objects.filter(

        # Only include consumers from the current cycle.
        bill_month_dt__gte=cycle_start,
        bill_month_dt__lte=cycle_end,

        # Only include consumers with a recorded reading date.
        reading_date_db__isnull=False,

    ).count()


    # ==================================================
    # CYCLE INFORMATION
    # ==================================================

    # Get the current day number within the month.
    #
    # Example:
    #
    # 11 September → day 11
    day_of_cycle = today.day


    # Get the total number of days in the current month.
    #
    # Example:
    #
    # September → 30 days
    # August    → 31 days
    cycle_days = cycle_end.day


    # Calculate the number of calendar days remaining
    # in the current reading cycle.
    #
    # Example:
    #
    # Cycle length → 30 days
    # Current day  → 11
    #
    # Remaining → 19 days
    #
    # Note:
    # This currently counts calendar days and does not
    # exclude weekends or holidays.
    working_left = cycle_days - day_of_cycle


    # ==================================================
    # ACTUAL PACE — MOST RECENT 7 DAYS
    # ==================================================

    # Calculate the date six days before today.
    #
    # Including today creates a seven-day window.
    #
    # Example:
    #
    # Today → 11 Sep
    #
    # Start → 05 Sep
    #
    # Range → 05 Sep to 11 Sep
    seven_days_ago = today - timedelta(days=6)


    # Count all readings completed during the most recent
    # seven-day period.
    readings_last_7_days = Consumers.objects.filter(

        # Include readings from the beginning of the
        # seven-day window.
        reading_date_db__gte=seven_days_ago,

        # Include readings up to today.
        reading_date_db__lte=today,

    ).count()


    # Calculate the average number of readings completed
    # per day during the last seven days.
    #
    # Example:
    #
    # 7,000 readings
    # ───────────── = 1,000 readings per day
    #      7
    actual_pace = round(
        readings_last_7_days / 7,
        2
    )


    # ==================================================
    # RETURN API RESPONSE
    # ==================================================

    # Return the calculated metrics as JSON.
    #
    # The property names intentionally match the data
    # structure expected by the React CycleStatus component.
    return JsonResponse({

        # Total completed readings in the current cycle.
        "total": total,

        # Total consumers expected in the current cycle.
        "expected": expected,

        # Current day of the cycle.
        "dayOfCycle": day_of_cycle,

        # Total days in the cycle.
        "cycleDays": cycle_days,

        # Remaining calendar days in the cycle.
        "workingLeft": working_left,

        # Average daily reading pace over the last 7 days.
        "actualPace": actual_pace,

    })


@parser_classes([MultiPartParser, FormParser])
@api_view(["POST"])

>>>>>>> b85f512 (Connect dashboard reading cycle to backend API)
def consumers(request):
    data = request.data.copy()
    # _mutable = data._mutable
    # data._mutable = True
    rdng_date = data["rdng_date"]
    cons_name = data["cons_name"]
    cons_ac_no = data["cons_ac_no"]
    ofc_section = data['ofc_section']

    if cons_name == "Test":
        return Response({"status": True, "message": "Test data not inserted"})

    reading_date_db = rdng_date[:10]
    adddate = "-01"
    bill_month_add = reading_date_db[:7] + adddate
    print(bill_month_add)

    reading_year = reading_date_db[:4]
    reading_month = reading_date_db[5:7]
    print(reading_date_db)
    print(reading_month)
    status = ""
    # comment this line
    data["reading_date_db"] = reading_date_db
    data["bill_month_dt"] = bill_month_add

    # NEW LOGIC : IF OCR READING IS NOT FOUND SET IMAGE BLUR
    if data.get("prsnt_ocr_rdng") == "Not Found":
        data["prsnt_rdng_ocr_excep"] = "Image blur"

    char = 0
    ba_bl_id = data["ba_bl_id"]
    try:
        if data["prsnt_mtr_status"] == "Ok":
            if data["prsnt_ocr_rdng"] != "Not Found":
                prsnt_ocr_rdng_temp = str(int(data["prsnt_ocr_rdng"]))
                prsnt_rdng_temp = str(int(data["prsnt_rdng"]))
                temp = max(int(prsnt_ocr_rdng_temp), int(prsnt_rdng_temp))
                temp1 = min(int(prsnt_ocr_rdng_temp), int(prsnt_rdng_temp))
                print("prsnt_ocr_rdng_temp--------->", prsnt_ocr_rdng_temp)
                print("prsnt_rdng_temp--------->", prsnt_rdng_temp)
                print("data['prsnt_ocr_rdng']--------->",
                      data["prsnt_ocr_rdng"])
                print("data['prsnt_rdng']--------->", data["prsnt_rdng"])
                flag = False
                if (prsnt_ocr_rdng_temp) == (prsnt_rdng_temp):
                    status = "Exact"

                elif (len(prsnt_ocr_rdng_temp) - len(prsnt_rdng_temp)) in (-1, 1):
                    # temp =max(int(prsnt_ocr_rdng_temp),int(prsnt_rdng_temp))
                    # temp1 =min(int(prsnt_ocr_rdng_temp),int(prsnt_rdng_temp))

                    for x in range(len(str(temp))):
                        temp_t = str(temp)
                        l = list(temp_t)
                        l.pop(x)
                        if l == list(str(temp1)):
                            flag = True
                    if flag == True:
                        status = "1_val_miss"
                    else:
                        status = "diff"
                elif len(prsnt_ocr_rdng_temp) == len(prsnt_rdng_temp):
                    for x in range(len(prsnt_ocr_rdng_temp)):
                        u, v = prsnt_ocr_rdng_temp, prsnt_rdng_temp
                        if u[x] != v[x]:
                            char += 1

                    if char == 1:
                        status = "1_val_diff"
                    else:
                        status = "diff"

                else:
                    print("OOKOOKOKOKOKOK")
                    if (str(temp1) in str(temp)) and (len(str(temp1)) > 1):
                        status = "subs"
                    else:
                        status = "diff"

                if status == "1_val_diff" or status == "1_val_miss" or status == "subs":
                    print("INSIDE UPDATESSSSS")
                    data["prsnt_rdng_ocr_odv"] = data["prsnt_ocr_rdng"]
                    data["rdng_ocr_status_odv"] = data["rdng_ocr_status"]
                    data["prsnt_ocr_excep_old_values"] = data["prsnt_rdng_ocr_excep"]
                    data["prsnt_ocr_rdng"] = data["prsnt_rdng"]
                    data["rdng_ocr_status"] = "Passed"
                    data["manual_update_flag"] = "true"
                    data["prsnt_rdng_ocr_excep"] = ""
                    data["rdng_ocr_status_changed_by"] = "Backend_CC"
    except:
        pass
    try:
        if data["prsnt_mtr_status"] == "Ok" and data["rdng_ocr_status"] == "":
            data["rdng_ocr_status"] = "Failed"
            data["manual_update_flag"] = "true"
            data["rdng_ocr_status_changed_by"] = "Backend_RERUN"
    except:
        pass

    # if bill id is present check the consumer ac no and month

    newid = (
        Consumers.objects.filter(
            Q(bill_month_dt=bill_month_add) & Q(
                cons_ac_no=cons_ac_no) & Q(ofc_section=ofc_section)
        )
        .order_by("-id")
        .first()
    )

    print("newid", newid)
    if newid is not None:
        # Check all the columns only then update
        if (
            newid.ofc_discom == data['ofc_discom'] and
            newid.ofc_zone == data['ofc_zone'] and
            newid.ofc_circle == data['ofc_circle'] and
            newid.ofc_division == data['ofc_division'] and
            newid.ofc_sub_div_code == data['ofc_sub_div_code'] and
            newid.ofc_subdivision == data['ofc_subdivision'] and
            newid.ofc_section == data['ofc_section'] and
            newid.mr_unit == data['mr_unit'] and
            newid.bl_area_code == data['bl_area_code'] and
            newid.bl_agnc_type == data['bl_agnc_type'] and
            newid.bl_agnc_name == data['bl_agnc_name'] and
            newid.mr_id == data['mr_id'] and
            newid.mr_ph_no == data['mr_ph_no'] and
            newid.cons_ac_no == data['cons_ac_no'] and
            newid.cons_name == data['cons_name'] and
            newid.con_trf_cat == data['con_trf_cat'] and
            newid.con_mtr_sl_no == data['con_mtr_sl_no'] and
            newid.con_mtr_phs == data['con_mtr_phs'] and
            newid.rdng_req_val == data['rdng_req_val'] and
            newid.prev_rdng == data['prev_rdng'] and
            newid.prev_md == data['prev_md'] and
            newid.prev_pf_rdng == data['prev_pf_rdng'] and
            newid.prev_rdng_date == data['prev_rdng_date'] and
            newid.prev_rdng_status == data['prev_rdng_status'] and
            newid.bl_mnth == data['bl_mnth'] and
            newid.rdng_date == data['rdng_date'] and
            newid.geo_lat == data['geo_lat'] and
            newid.geo_long == data['geo_long'] and
            newid.prsnt_mtr_status == data['prsnt_mtr_status'] and
            newid.abnormality == data['abnormality'] and
            newid.mr_rmrk == data['mr_rmrk'] and
            newid.rdng_ocr_status == data['rdng_ocr_status'] and
            newid.prsnt_ocr_rdng == data['prsnt_ocr_rdng'] and
            newid.prsnt_rdng == data['prsnt_rdng'] and
            newid.prsnt_rdng_ocr_excep == data['prsnt_rdng_ocr_excep'] and
            newid.rdng_img == data['rdng_img'] and
            newid.ocr_md_status == data['ocr_md_status'] and
            newid.prsnt_md_rdng_ocr == data['prsnt_md_rdng_ocr'] and
            newid.prsnt_md_rdng == data['prsnt_md_rdng'] and
            newid.md_ocr_excep == data['md_ocr_excep'] and
            newid.md_img == data['md_img'] and
            newid.ocr_pf_status == data['ocr_pf_status'] and
            newid.ocr_pf_reading == data['ocr_pf_reading'] and
            newid.pf_manual_reading == data['pf_manual_reading'] and
            newid.pf_ocr_exception == data['pf_ocr_exception'] and
            newid.pf_image == data['pf_image'] and
            newid.ai_mdl_ver == data['ai_mdl_ver'] and
            newid.ph_name == data['ph_name'] and
            newid.cmra_res == data['cmra_res'] and
            newid.andr_ver == data['andr_ver'] and
            newid.qc_req == data['qc_req'] and
            newid.ba_cons_id == data['ba_cons_id'] and
            newid.ba_ac_id == data['ba_ac_id'] and
            newid.ba_prsnt_rdng_status == data['ba_prsnt_rdng_status'] and
            newid.ba_mrc == data['ba_mrc'] and
            newid.ba_mru == data['ba_mru'] and
            newid.ba_subdiv == data['ba_subdiv'] and
            newid.ba_div == data['ba_div'] and
            newid.ba_agnc_id == data['ba_agnc_id'] and
            newid.ba_bl_id == data['ba_bl_id'] and
            newid.ba_bl_date == data['ba_bl_date'] and
            newid.ba_prev_rdng_status == data['ba_prev_rdng_status'] and
            newid.qc_done == data['qc_done'] and
            newid.qc_done_user_id == data['qc_done_user_id'] and
            newid.qc_date == data['qc_date'] and
            newid.qc_flag == data['qc_flag'] and
            newid.qc_rmrk == data['qc_rmrk'] and
            newid.ai_retrain == data['ai_retrain'] and
            newid.is_object_meter == data['is_object_meter'] and
            newid.mr_success_feedback == data['mr_success_feedback'] and
            newid.reading_parameter_type == data['reading_parameter_type'] and
            newid.md_reading_parameter_type == data['md_reading_parameter_type'] and
            newid.pf_reading_parameter_type == data['pf_reading_parameter_type'] and
            newid.rdng_ocr_status_changed_by == data['rdng_ocr_status_changed_by'] and
            newid.prsnt_rdng_ocr_odv == data['prsnt_rdng_ocr_odv'] and
            newid.rdng_ocr_status_odv == data['rdng_ocr_status_odv'] and
            newid.prsnt_ocr_excep_old_values == data['prsnt_ocr_excep_old_values']
        ):
            print("updatation does not takes place")
            return Response({"status": True, "message": "No change in Data"})
        else:
            print("updating as some of the values are changed")

            serializer = ConsumerSerializer(newid, data=data, partial=True)
            if serializer.is_valid():
                serializer.save()
                return Response({"status": True, "message": "Data Updated successfully"})
            return Response(serializer.errors)

    # if bill id is not present then insert
    print("inserted")
    serializer = ConsumerSerializer(data=data)
    if serializer.is_valid():
        serializer.save()
        return Response(
            {"status": True, "message": "Data added successfully", "version": "28"}
        )
    return Response(serializer.errors)


@api_view(["POST"])
def consumers_bulk(request):
    data_list = request.data.copy()
    print("datalist--->", data_list)
    count_insert = 0
    count_update = 0
    failed_consumers = []
    for data in data_list:
        # ensure defaults for new keys
        data['kvah_manual'] = data.get('kvah_manual', None)
        data['kvah_Status'] = data.get('kvah_Status', None)
        data['mtr_sr_no'] = data.get('mtr_sr_no', None)
        data['ekwh_img'] = data.get('ekwh_img', None)
        data['ekwh_ocr_rdng'] = data.get('ekwh_ocr_rdng', None)
        data['ekwh_manual_rdng'] = data.get('ekwh_manual_rdng', None)
        data['ekvah_img'] = data.get('ekvah_img', None)
        data['ekvah_ocr_rdng'] = data.get('ekvah_ocr_rdng', None)
        data['ekvah_manual_rdng'] = data.get('ekvah_manual_rdng', None)

        #----------spoofed_status---------------

        value = data.get('is_spoofed', False)

        if isinstance(value, str):
            data['is_spoofed'] = value.strip().lower() in ['true', '1', 'yes']
        elif isinstance(value, (int, float)):
            data['is_spoofed'] = value == 1
        elif isinstance(value, bool):
            data['is_spoofed'] = value
        else:
            data['is_spoofed'] = False

        if data['is_spoofed']:
            data['prsnt_rdng_ocr_excep'] = 'Spoofed Image'


        # -----

        rdng_date = data["rdng_date"]
        cons_name = data["cons_name"]
        cons_ac_no = data["cons_ac_no"]
        ofc_section = data['ofc_section']

        if cons_name == "Test":
            return Response({"status": True, "message": "Test data not inserted"})

        reading_date_db = rdng_date[:10]
        adddate = "-01"
        bill_month_add = reading_date_db[:7] + adddate
        print(bill_month_add)
 
        reading_year = reading_date_db[:4]
        reading_month = reading_date_db[5:7]
        print(reading_date_db)
        print(reading_month)
        status = ""
        # comment this line
        data["reading_date_db"] = reading_date_db
        data["bill_month_dt"] = bill_month_add


        char = 0
        ba_bl_id = data["ba_bl_id"]

        rdngImg = data.get("rdng_img")

        # Ensure rdngImg is a clean string (not list or nested)
        if isinstance(rdngImg, list):
            rdngImg = rdngImg[0]
        elif isinstance(rdngImg, dict) and "url" in rdngImg:
            rdngImg = rdngImg["url"]
        elif not isinstance(rdngImg, str):
            rdngImg = str(rdngImg)

        # strip extra quotes or spaces
        rdngImg = rdngImg.strip('"').strip()

        try:
            if data["prsnt_mtr_status"] == "Ok":
                if data["prsnt_ocr_rdng"] != "Not Found":
                    prsnt_ocr_rdng_temp = str(int(data["prsnt_ocr_rdng"]))
                    prsnt_rdng_temp = str(int(data["prsnt_rdng"]))
                    temp = max(int(prsnt_ocr_rdng_temp), int(prsnt_rdng_temp))
                    temp1 = min(int(prsnt_ocr_rdng_temp), int(prsnt_rdng_temp))
                    print("prsnt_ocr_rdng_temp--------->", prsnt_ocr_rdng_temp)
                    print("prsnt_rdng_temp--------->", prsnt_rdng_temp)
                    print("data['prsnt_ocr_rdng']--------->",
                          data["prsnt_ocr_rdng"])
                    print("data['prsnt_rdng']--------->", data["prsnt_rdng"])
                    flag = False
                    if (prsnt_ocr_rdng_temp) == (prsnt_rdng_temp):
                        status = "Exact"
 
                    elif (len(prsnt_ocr_rdng_temp) - len(prsnt_rdng_temp)) in (-1, 1):
                        for x in range(len(str(temp))):
                            temp_t = str(temp)
                            l = list(temp_t)
                            l.pop(x)
                            if l == list(str(temp1)):
                                flag = True
                        if flag == True:
                            status = "1_val_miss"
                        else:
                            status = "diff"
                    elif len(prsnt_ocr_rdng_temp) == len(prsnt_rdng_temp):
                        for x in range(len(prsnt_ocr_rdng_temp)):
                            u, v = prsnt_ocr_rdng_temp, prsnt_rdng_temp
                            if u[x] != v[x]:
                                char += 1
 
                        if char == 1:
                            status = "1_val_diff"
                        else:
                            status = "diff"
 
                    else:
                        print("OOKOOKOKOKOKOK")
                        if (str(temp1) in str(temp)) and (len(str(temp1)) > 1):
                            status = "subs"
                        else:
                            status = "diff"
 
                    if status == "1_val_diff" or status == "1_val_miss" or status == "subs":
                        print("INSIDE UPDATESSSSS")
                        data["prsnt_rdng_ocr_odv"] = data["prsnt_ocr_rdng"]
                        data["rdng_ocr_status_odv"] = data["rdng_ocr_status"]
                        data["prsnt_ocr_excep_old_values"] = data["prsnt_rdng_ocr_excep"]
                        data["prsnt_ocr_rdng"] = data["prsnt_rdng"]
                        data["rdng_ocr_status"] = "Passed"
                        data["manual_update_flag"] = "true"
                        data["prsnt_rdng_ocr_excep"] = ""
                        data["rdng_ocr_status_changed_by"] = "Backend_CC"
        except:
            pass
        try:
            if data["prsnt_mtr_status"] == "Ok" and data["rdng_ocr_status"] == "" :
                data["rdng_ocr_status"] = "Failed"
                data["manual_update_flag"] = "true"
                data["rdng_ocr_status_changed_by"] = "Backend_RERUN"
        except:
            pass
        try:
            newid = (
                Consumers.objects.filter(
                    Q(bill_month_dt=bill_month_add) & Q(
                        cons_ac_no=cons_ac_no) & Q(ofc_section=ofc_section)
                )
                .order_by("-id")
                .first()
            )
            print("newid", newid)
            if newid is not None:
                # Check all the columns only then update
                if (
                    newid.ofc_discom == data['ofc_discom'] and
                    newid.ofc_zone == data['ofc_zone'] and
                    newid.ofc_circle == data['ofc_circle'] and
                    newid.ofc_division == data['ofc_division'] and
                    newid.ofc_sub_div_code == data['ofc_sub_div_code'] and
                    newid.ofc_subdivision == data['ofc_subdivision'] and
                    newid.ofc_section == data['ofc_section'] and
                    newid.mr_unit == data['mr_unit'] and
                    newid.bl_area_code == data['bl_area_code'] and
                    newid.bl_agnc_type == data['bl_agnc_type'] and
                    newid.bl_agnc_name == data['bl_agnc_name'] and
                    newid.mr_id == data['mr_id'] and
                    newid.mr_ph_no == data['mr_ph_no'] and
                    newid.cons_ac_no == data['cons_ac_no'] and
                    newid.cons_name == data['cons_name'] and
                    newid.con_trf_cat == data['con_trf_cat'] and
                    newid.con_mtr_sl_no == data['con_mtr_sl_no'] and
                    newid.con_mtr_phs == data['con_mtr_phs'] and
                    newid.rdng_req_val == data['rdng_req_val'] and
                    newid.prev_rdng == data['prev_rdng'] and
                    newid.prev_md == data['prev_md'] and
                    newid.prev_pf_rdng == data['prev_pf_rdng'] and
                    newid.prev_rdng_date == data['prev_rdng_date'] and
                    newid.prev_rdng_status == data['prev_rdng_status'] and
                    newid.bl_mnth == data['bl_mnth'] and
                    newid.rdng_date == data['rdng_date'] and
                    newid.geo_lat == data['geo_lat'] and
                    newid.geo_long == data['geo_long'] and
                    newid.prsnt_mtr_status == data['prsnt_mtr_status'] and
                    newid.abnormality == data['abnormality'] and
                    newid.mr_rmrk == data['mr_rmrk'] and
                    newid.rdng_ocr_status == data['rdng_ocr_status'] and
                    newid.prsnt_ocr_rdng == data['prsnt_ocr_rdng'] and
                    newid.prsnt_rdng == data['prsnt_rdng'] and
                    newid.prsnt_rdng_ocr_excep == data['prsnt_rdng_ocr_excep'] and
                    newid.rdng_img == data['rdng_img'] and
                    newid.ocr_md_status == data['ocr_md_status'] and
                    newid.prsnt_md_rdng_ocr == data['prsnt_md_rdng_ocr'] and
                    newid.prsnt_md_rdng == data['prsnt_md_rdng'] and
                    newid.md_ocr_excep == data['md_ocr_excep'] and
                    newid.md_img == data['md_img'] and
                    newid.ocr_pf_status == data['ocr_pf_status'] and
                    newid.ocr_pf_reading == data['ocr_pf_reading'] and
                    newid.pf_manual_reading == data['pf_manual_reading'] and
                    newid.pf_ocr_exception == data['pf_ocr_exception'] and
                    newid.pf_image == data['pf_image'] and
                    newid.ai_mdl_ver == data['ai_mdl_ver'] and
                    newid.ph_name == data['ph_name'] and
                    newid.cmra_res == data['cmra_res'] and
                    newid.andr_ver == data['andr_ver'] and
                    newid.data_sync_date == data['data_sync_date'] and
                    newid.qc_req == data['qc_req'] and
                    newid.ba_cons_id == data['ba_cons_id'] and
                    newid.ba_ac_id == data['ba_ac_id'] and
                    newid.ba_prsnt_rdng_status == data['ba_prsnt_rdng_status'] and
                    newid.ba_mrc == data['ba_mrc'] and
                    newid.ba_mru == data['ba_mru'] and
                    newid.ba_subdiv == data['ba_subdiv'] and
                    newid.ba_div == data['ba_div'] and
                    newid.ba_agnc_id == data['ba_agnc_id'] and
                    newid.ba_bl_id == data['ba_bl_id'] and
                    newid.ba_bl_date == data['ba_bl_date'] and
                    newid.ba_prev_rdng_status == data['ba_prev_rdng_status'] and
                    newid.qc_done == data['qc_done'] and
                    newid.qc_done_user_id == data['qc_done_user_id'] and
                    newid.qc_date == data['qc_date'] and
                    newid.qc_flag == data['qc_flag'] and
                    newid.qc_rmrk == data['qc_rmrk'] and
                    newid.ai_retrain == data['ai_retrain'] and
                    newid.is_object_meter == data['is_object_meter'] and
                    newid.mr_success_feedback == data['mr_success_feedback'] and
                    newid.reading_parameter_type == data['reading_parameter_type'] and
                    newid.md_reading_parameter_type == data['md_reading_parameter_type'] and
                    newid.pf_reading_parameter_type == data['pf_reading_parameter_type'] and
                    newid.rdng_ocr_status_changed_by == data['rdng_ocr_status_changed_by'] and
                    newid.prsnt_rdng_ocr_odv == data['prsnt_rdng_ocr_odv'] and
                    newid.rdng_ocr_status_odv == data['rdng_ocr_status_odv'] and
                    newid.prsnt_ocr_excep_old_values == data['prsnt_ocr_excep_old_values'] and
                    newid.kvah_rdng == data['kvah_rdng'] and
                    newid.kvah_img == data['kvah_img'] and
                    newid.kvah_manual == data['kvah_manual'] and
                    newid.kvah_Status == data['kvah_Status'] and
                    newid.mtr_sr_no == data['mtr_sr_no'] and 
                    newid.is_spoofed == data['is_spoofed'] and
                    newid.ekwh_img == data['ekwh_img'] and
                    newid.ekwh_ocr_rdng == data['ekwh_ocr_rdng'] and
                    newid.ekwh_manual_rdng == data['ekwh_manual_rdng'] and
                    newid.ekvah_img == data['ekvah_img'] and
                    newid.ekvah_ocr_rdng == data['ekvah_ocr_rdng'] and
                    newid.ekvah_manual_rdng == data['ekvah_manual_rdng'] and
                    newid.calculatedkvah == data['calculatedkvah'] and
                    newid.calculatedkva == data['calculatedkva']
                ):
                    print("updatation does not takes place")
                    # return Response({"status": True, "message": "No change in Data"})
                else:
                    print("updating as some of the values are changed")

                    serializer = ConsumerSerializer(
                        newid, data=data, partial=True)
                    if serializer.is_valid():
                        serializer.save()
                        count_update += 1

            # if bill id is not present then insert
 
            else:
                print("inserted")
                serializer = ConsumerSerializer(data=data)
                if serializer.is_valid():
                    serializer.save()
                    count_insert += 1

        except Exception as e:
            failedCons = {"message": str(e), "cons_ac_no": data['cons_ac_no']}
            failed_consumers.append(failedCons)
    print({"status": True, "message": f"Data inserted {count_insert} and Data updated {count_update}", "version": "28"})
    if len(failed_consumers) > 0:
        return Response(
            {"status": False, "message": f"Data inserted {count_insert} and Data updated {count_update}",
                "failed_consumers": failed_consumers, "version": "28"}
        )
    return Response(
        {"status": True, "message": f"Data inserted {count_insert} and Data updated {count_update}",
            "failed_consumers": failed_consumers, "version": "28"}
    )


from rest_framework.decorators import api_view
from rest_framework.response import Response
from django.db.models import Q
from .models import Consumers


@api_view(["POST"])
def consumerscheck(request):

    try:
        cons_ids = request.data.get("cons_ids", [])
        start_date = request.data.get("start_date")
        end_date = request.data.get("end_date")

        if not cons_ids:
            return Response({
                "status": False,
                "message": "cons_ids list required"
            })

        # ---- SINGLE DB QUERY (optimized) ----
        existing_ids = set(
            Consumers.objects.filter(
                cons_ac_no__in=cons_ids,
                rdng_date__range=[start_date, end_date]
            ).values_list("cons_ac_no", flat=True)
        )

        # ---- FIND FAILED IDS ----
        input_ids = set(cons_ids)
        failed_ids = list(input_ids - existing_ids)

        return Response({
            "status": True,
            "total_requested": len(cons_ids),
            "found_in_db": len(existing_ids),
            "failed_ids": failed_ids
        })

    except Exception as e:
        return Response({
            "status": False,
            "error": str(e)
        })
    
@api_view(["GET"])
def getconsumers(request):
    data = Consumers.objects.all().order_by("-id")
    serializer = ConsumerSerializer(data, many=True)
    return Response(serializer.data)


@api_view(["POST"])
def deleteconsumers(request):
    try:
        mrId = request.data.get("mrId")
        user = MeterReaderRegistration.objects.get(mrId=mrId)
        user.delete()
        return Response(
            {"status": True, "message": "MR deleted successfully(from api)"}
        )
    except MeterReaderRegistration.DoesNotExist:
        return Response(
            {"status": False, "message": "MR Does not Exist(from api)"},
            status=status.HTTP_200_OK,
        )


@api_view(["GET"])
def getofficedatalist(request):
    data = Office.objects.all()
    cursor = connection.cursor()
    query = f"""
    select * from office

    """
    cursor.execute(query)
    result = dictfetchall(cursor)
    return Response(result)
    pass


@parser_classes([MultiPartParser, FormParser])
@api_view(["POST"])
def meterReaderRegistrationfun(request):
    data = request.data
    serializer = MeterReaderRegistrationSerializer(data=data)
    if serializer.is_valid():
        serializer.save()
        return Response(
            {"status": True, "message": "Registration Successfull (from api)"},
            status=status.HTTP_201_CREATED,
        )
    return Response(
        {"status": False,
            "message": "Meter Reader Id Already exists (from api)"},
        status=status.HTTP_200_OK,
    )


@parser_classes([MultiPartParser, FormParser])
@api_view(["POST"])
def meterReaderRegistrationUpdateOffice(request):
    data = request.data
    mrId = data["mrId"]
    id = MeterReaderRegistration.objects.get(mrId=mrId)
    if id is None:
        return Response(
            {"status": False,
                "message": "Meter Reader Id Does Not exists (from api)"},
            status=status.HTTP_200_OK,
        )
    serializer = MeterReaderRegistrationSerializer(id, data=data, partial=True)
    if serializer.is_valid():
        serializer.save()
        return Response(
            {
                "status": True,
                "message": "Registration Data Updated Successfull (from api)",
            },
            status=status.HTTP_201_CREATED,
        )
    return Response(
        {"status": False, "message": "Error (from api)"}, status=status.HTTP_200_OK
    )


@parser_classes([MultiPartParser, FormParser])
@api_view(["POST"])
def metereReaderlogin(request):
    newdata = request.data
    mrId = newdata["mrId"]
    token = newdata["androidToken"]
    try:
        id = MeterReaderRegistration.objects.get(mrId=mrId)
        serializer = MeterReaderRegistrationSerializer(
            id, data=newdata, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(
                {
                    "status": True,
                    "message": "Login Successfull(from api)",
                    "data": serializer.data,
                },
                status=status.HTTP_201_CREATED,
            )
        return Response(
            {"status": False, "message": "Something wrong(from api)"},
            status=status.HTTP_200_OK,
        )
    except MeterReaderRegistration.DoesNotExist:
        return Response(
            {"status": False, "message": "MR Does not Exist(from api)"},
            status=status.HTTP_200_OK,
        )

from .models import (
    MeterReaderRegistration,
    SupervisorLogin,
    SupervsiorLocation,
    Consumers,
)

from .serializers import (
    MeterReaderRegistrationSerializer,
    SupervisorLoginSerializer,
)


from datetime import date
from django.db.models import Exists, OuterRef, Q
from rest_framework.decorators import api_view
from rest_framework.response import Response

@api_view(["GET"])
def getregdata(request):
    role_to_fetch = request.query_params.get("role", "meterreader").lower()
    discom = request.query_params.get("discom", "all").upper()

    from_date = request.query_params.get("from_date")
    to_date = request.query_params.get("to_date")

    print("ROLE:", role_to_fetch)
    print("DISCOM:", discom)
    print("FROM:", from_date)
    print("TO:", to_date)

    if role_to_fetch == "supervisor":

        today = date.today()

        location_exists = SupervsiorLocation.objects.filter(
            supervisor_number=OuterRef("supervisor_number"),
            date=today
        )

        filters = {}

        if discom != "ALL":
            filters["discom"] = discom

        qs = (
            SupervisorLogin.objects
            .filter(**filters)
            .filter(
                id__in=SupervisorLogin.objects.values("supervisor_number")
                .distinct()
                .values_list("id", flat=True)
            )
            .annotate(location=Exists(location_exists))
            .order_by("supervisor_number", "-location")
            .distinct("supervisor_number")
        )

        result = list(qs)
        result.sort(key=lambda x: not x.location)

        serializer = SupervisorLoginSerializer(result, many=True)
        return Response(serializer.data)

    else:

        data = MeterReaderRegistration.objects.all()

        if discom != "ALL":
            data = data.filter(discom__iexact=discom)

        if from_date and to_date:

            # equivalent to SQL BETWEEN from_date AND to_date
            mr_ids = (
                Consumers.objects
                .filter(
                    reading_date_db__gte=from_date,
                    reading_date_db__lte=to_date
                )
                .exclude(mr_id__isnull=True)
                .exclude(mr_id="")
                .values_list("mr_id", flat=True)
                .distinct()
            )

            print("MR COUNT IN READINGMASTER:", mr_ids.count())
            print("FIRST 20 MRS:", list(mr_ids[:20]))

            data = data.filter(
                mrId__in=mr_ids
            )

        print("FINAL MR COUNT:", data.count())

        serializer = MeterReaderRegistrationSerializer(data, many=True)

        return Response(serializer.data)



@api_view(["GET"])
def get_officedata(request):
    with open("D:/SBPDCL.json") as f:
        data = json.load(f)
        for i in data:
            data = Office.objects.create(
                discom=i["discom"],
                zone=i["zone"],
                circlename=i["circlename"],
                divisionname=i["divisionname"],
                divisioncode=i["divisioncode"],
                subdivision=i["subdivision"],
                subdivisioncode=i["subdivisioncode"],
                sectionname=i["sectionname"],
                sectioncode=i["sectioncode"],
            )
            data.save()
    return Response("data added Successfully")


# Create your views here.
@api_view(["POST"])
def loginuser(request):
    data = request.data
    email = data["email"]
    password = data["password"]
    print("email", email)
    print("password", password)
    try:
        if email == "payfinix@gmail.com" and password == "payfinix#123":
            print("hello")
            token = jwt.encode({"agentid": email},
                               SECRETKEY, algorithm="HS256")
            # return Response({"user":"admin","token":token})
            return Response({"user": "admin", "accessToken": token})

        return Response(
            {"status": "False", "Msg": "Email or Password did not match"},
            status=status.HTTP_404_NOT_FOUND,
        )
    except:
        return Response(
            {"status": "False", "Msg": "Email or Password did not match"},
            status=status.HTTP_404_NOT_FOUND,
        )


@api_view(["GET"])
def get_discom(request):
    newlist = []
    discom = Office.objects.values_list("discom").distinct()
    for row in discom:
        print(row[0])
        newlist.append(row[0])
        print(newlist)
    return Response(newlist)


@api_view(["POST"])
def get_zone(request):
    newlist = []
    zone = (
        Office.objects.filter(discom=request.data["discom"])
        .values_list("zone")
        .distinct()
    )
    for row in zone:
        print(row[0])
        newlist.append(row[0])
        print(newlist)
    return Response(newlist)


@api_view(["POST"])
def get_circle(request):
    newlist = []
    circle = (
        Office.objects.filter(zone=request.data["zone"])
        .values_list("circlename")
        .distinct()
    )
    for row in circle:
        print(row[0])
        newlist.append(row[0])
        print(newlist)
    return Response(newlist)


@api_view(["POST"])
def get_division(request):
    newlist = []
    division = (
        Office.objects.filter(circlename=request.data["circle"])
        .values_list("divisionname")
        .distinct()
    )
    for row in division:
        print(row[0])
        newlist.append(row[0])
        print(newlist)
    return Response(newlist)


@api_view(["POST"])
def get_subdivision(request):
    newlist = []
    subdivision = (
        Office.objects.filter(divisionname=request.data["division"])
        .values_list("subdivision")
        .distinct()
    )
    for row in subdivision:
        print(row[0])
        newlist.append(row[0])
        print(newlist)
    return Response(newlist)


@api_view(["POST"])
def get_section(request):
    new = []

    def listfun(dict):
        print(dict)
        new.append(dict.copy())
        return new

    newdict = {}
    section = (
        Office.objects.filter(subdivision=request.data["subdivision"])
        .values_list("sectionname", "sectioncode")
        .distinct()
    )

    print("section", section)
    for row in section:
        newdict["sectionName"] = row[0]
        newdict["sectionCode"] = row[1]

        newdata = listfun(newdict)
    return Response(newdata)


@api_view(["POST"])
def get_sectionforuser(request):
    newlist = []
    section = (
        Office.objects.filter(subdivision=request.data["subdivision"])
        .values_list("sectionname")
        .distinct()
    )
    print("section", section)
    for row in section:
        newlist.append(row[0])
    return Response(newlist)



@api_view(["GET"])
def get_meter_reader_detail(request):
    pagesize = request.query_params.get("pagesize", 10)
    paginator = PageNumberPagination()
    paginator.page_size = pagesize
    monthByData = request.query_params.get("getMonth", None)
    mrid = request.query_params.get("mrid", None)
    startDate = request.query_params.get("startdate", None)
    endDate = request.query_params.get("enddate", None)
    person_objects = []
    if monthByData:
        monthByData = monthByData.split("-")[1]
        person_objects = Consumers.objects.filter(
            reading_date_db__month=monthByData
        ).order_by("-id")
    elif startDate and endDate:
        person_objects = Consumers.objects.filter(
            reading_date_db__range=[startDate, endDate]
        ).order_by("-id")
    elif mrid:
        person_objects = Consumers.objects.filter(mr_id=mrid).order_by("-id")
    else:
        person_objects = []
    result_page = paginator.paginate_queryset(person_objects, request)
    serializer = ConsumerDataSerializer(result_page, many=True)
    return paginator.get_paginated_response(serializer.data)




@api_view(["GET"])
def consumer_wise_details(request):
    acno = request.query_params.get("acno", None)
    print(acno)
    if acno is not None:
        data = Consumers.objects.filter(cons_ac_no=acno)
        serializer = ConsumerWiseDetailsSerializer(data, many=True)
        return Response(serializer.data)
    return Response({"status": False, "consumer": "Not Available"})


@api_view(["GET"])
def getconsumerscount(request):
    uniqueConsAcc = Consumers.objects.values_list("cons_ac_no").distinct()
    consumerscount = len(uniqueConsAcc)
    return Response({"Consumers": consumerscount})

@api_view(["POST"])
def getmridforSection(request):
    data = request.data.get("sectioncode", None)
    newlist = []
    mridData = (
        Consumers.objects.filter(
            ofc_section=data).values_list("mr_id").distinct()
    )
    # serializer = ConsumerDataSerializer(mridData,many=True)
    for row in mridData:
        print(row[0])
        newlist.append(row[0])
        print(newlist)
    return Response(newlist)


@api_view(["GET"])
def test(request):
    offset = request.query_params.get("offset", None)
    new = []

    def listfun(dict):
        new.append(dict.copy())
        return new

    cursor = connection.cursor()
    new_dict = {}
    if offset is not None:
        clause = "OFFSET %s"
        params = [offset]
    else:
        clause = ""
        params = ""

    query = f"""select m.mr_id,m.rdng_date,m.prsnt_mtr_status,m.prsnt_ocr_rdng,m.prsnt_rdng,m.ocr_pf_reading,m.cons_name,m.prsnt_md_rdng_ocr,m.rdng_ocr_status,m.rdng_img,m.prsnt_md_rdng,m.id,r."mrPhoto"

                    from readingmaster m,meterreaderregistration r where m.mr_id=r."mrId" ORDER BY m.rdng_date DESC LIMIT 10 {clause}
    """

    cursor.execute(query, params)
    result = cursor.fetchall()
    print("RESULTS____________", result)
    try:
        for row in result:
            new_dict["id"] = row[11]
            new_dict["mrid"] = row[0]
            new_dict["rdngDate"] = row[1]
            new_dict["prsntmtrstatus"] = row[2]
            new_dict["prsntOcrRdng"] = row[3]
            new_dict["prsntRdng"] = row[4]
            new_dict["prsntPf"] = row[5]
            new_dict["consName"] = row[6]
            new_dict["prsntMdRdngOcr"] = row[7]
            new_dict["rdngocrstatus"] = row[8]
            new_dict["rdngImg"] = row[9]
            new_dict["prsntMdRdng"] = row[10]
            new_dict["avatar"] = row[12]
            newdata = listfun(new_dict)
        return Response(newdata)
    except:
        return Response([])


def dictfetchall(cursor):
    columns = [col[0] for col in cursor.description]
    return [dict(zip(columns, row)) for row in cursor.fetchall()]


@api_view(["GET"])
def consdetail(request):
    print("hi")
    acno = request.query_params.get("acno", None)
    cursor = connection.cursor()

    cursor.execute(
        f"""
    SELECT  m.ofc_discom,m.ofc_zone,m.ofc_circle,m.ofc_division,m.ofc_subdivision,m.ofc_section,
    m.id,m.cons_name,m.cons_ac_no,cons_address,m.cons_ph_no,m.con_trf_cat,m.mr_unit,
    r."mrId",r."mrName",r."mrPhone",r."mrPhoto" as avatar,m.con_mtr_sl_no,
    m.rdng_date,m.prsnt_mtr_status,m.prsnt_md_rdng,m.ocr_pf_reading,m.abnormality,m.prsnt_rdng_ocr_excep,m.md_ocr_excep,m.mr_rmrk,m.qc_req,m.ai_mdl_ver,m.ph_name,m.cmra_res,m.andr_ver,m.reading_date_db,
    m.rdng_img,m.md_img,m.pf_image,m.prsnt_ocr_rdng,m.prsnt_rdng,m.prsnt_md_rdng_ocr,m.rdng_ocr_status,m.ocr_pf_status,m.pf_manual_reading,m.prsnt_md_rdng,m.ocr_md_status,m.md_img
      FROM
    readingmaster m,meterreaderregistration r where m.mr_id=r."mrId" AND m.cons_ac_no='{acno}'

    """
    )
    results = dictfetchall(cursor)

    return Response(results)

@api_view(
    [
        "GET",
    ]
)




@parser_classes([MultiPartParser, FormParser])
@api_view(["POST"])
def consumerstest(request):
    data = request.data
    _mutable = data._mutable
    data._mutable = True
    rdng_date = data["rdng_date"]
    print("rdng_date", rdng_date)
    reading_date_db = rdng_date[:10]
    print("reading_date_db", reading_date_db)
    data["reading_date_db"] = reading_date_db
    data._mutable = _mutable

    return Response("ok")


@api_view(["GET"])
def clusters(request):
    # paginator = PageNumberPagination()
    # paginator.page_size = 100
    today = date.today()

    # data = Consumers.objects.all()

    # serializer = ConsumerSerializer(data, many=True)
    # result_page = paginator.paginate_queryset(serializer.data, request)
    # # serializer = ConsumerDataSerializer1(result_page, many=True)
    # clause = ''
    cursor = connection.cursor()
    query = f"""
    select mr_id,rdng_date,cons_name,geo_lat,geo_long,prsnt_mtr_status,rdng_ocr_status,rdng_img from readingmaster where reading_date_db='{today}'
    """
    cursor.execute(query)
    result = dictfetchall(cursor)
    return Response(result)

@api_view(["GET"])
def testdata(request):
    todaydate = date.today()
    new = []

    def listfun(dict):
        print(dict)
        new.append(dict.copy())
        return new

    newdict = {}
    cursor = connection.cursor()
    query = f"""SELECT r.mr_id,r.ofc_discom,r.ofc_zone,r.ofc_circle,r.ofc_division,r.ofc_subdivision,count(r.cons_ac_no) as billed_consumers,
count(r.prsnt_mtr_status='Ok' OR NULL) as ok_readings,count(r.rdng_ocr_status='Passed' or NULL) as OCRwithoutException,
count(r.rdng_ocr_status='Failed' OR NULL) as OCRwithException, m."mrPhone"
FROM readingmaster r
JOIN meterreaderregistration m ON m."mrId" = r.mr_id
AND reading_date_db ='{todaydate}'
GROUP BY r.mr_id,r.ofc_discom,r.ofc_zone,r.ofc_circle,r.ofc_division,r.ofc_subdivision, m."mrPhone";
        """
    cursor.execute(query)
    results = cursor.fetchall()
    try:
        for row in results:
            total = row[6]
            okreadings = row[7]
            ocrreadings = row[8]
            ocrwithexcep = row[9]

            # Percentage
            okreadpercent = round(((okreadings / total) * 100), 2)
            ocrreadingpercent = round(
                (((ocrreadings / okreadings) if okreadings else 0) * 100), 2
            )
            ocrwithexceppercent = round(
                (((ocrwithexcep / okreadings) if okreadings else 0) * 100), 2
            )

            # add to dictionary
            newdict["mrid"] = row[0]
            newdict["mrPhone"] = row[10]
            newdict["ofc_discom"] = row[1]
            newdict["ofc_zone"] = row[2]
            newdict["ofc_circle"] = row[3]
            newdict["ofc_division"] = row[4]
            newdict["ofc_subdivision"] = row[5]
            newdict["billed_consumers"] = row[6]

            newdict["OKreadings"] = okreadings
            newdict["OCRReadings"] = ocrreadings
            newdict["OCRwithException"] = ocrwithexcep

            newdict["OKreadingspercent"] = okreadpercent
            newdict["OCRReadingspercent"] = ocrreadingpercent
            newdict["OCRwithExceptionpercent"] = ocrwithexceppercent

            # add to list
            newdata = listfun(newdict)
        print("newdata--", newdata)
        wb = Workbook()
        ws = wb.active
        ws.title = "JSON Data"
        headers = list(newdata[0].keys())
        for j, header in enumerate(headers):
            ws.cell(row=1, column=j + 1, value=header)

        for i, row in enumerate(newdata, start=2):
            for j, key in enumerate(headers):
                ws.cell(row=i, column=j + 1, value=row[key])
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = "attachment; filename=data.xlsx"
        wb.save(response)
        return response

    except:
        return Response([])


@api_view(["GET"])
def geocluster(request):
    today = date.today()
    print(today)
    cursor = connection.cursor()
    query = f"""
    select m.mr_id as id,m.geo_lat,m.geo_long from readingmaster m where reading_date_db='{today}'

    """

    cursor.execute(query)

    result = dictfetchall(cursor)
    geo = to_geojson(result)

    return Response(geo)

    pass


def convertdicttolist(lsts, key):
    return [x.get(key) for x in lsts]


@api_view(["POST"])
def locationzone(request):
    locationwise = request.data.get("locationwise")
    locationname = request.data.get("locationname")
    groupby = request.data.get("groupby")
    where = request.data.get("where")
    previouslocation = request.data.get("previouslocation")
    clause = ""
    print("groupby", groupby)
    cursor = connection.cursor()
    if (locationwise is not None) and (locationname == "all"):
        location = "ofc_" + locationwise
        clause = "WHERE " + previouslocation + "='" + where + "' "
        cursor.execute(
            f"""
   select {groupby} as location
    from office {clause}  GROUP BY {groupby}
    """
        )
    if (locationwise is not None) and (locationname != "all"):
        clause = "WHERE " + locationwise + "='" + locationname + "' "
        cursor.execute(
            f"""
   select {groupby} as location
    from office {clause}  GROUP BY {groupby}
    """
        )

    result = dictfetchall(cursor)
    res = convertdicttolist(result, "location")
    return Response(res)


@api_view(["POST"])
def locationwisezone(request):
    month = date.today().month

    newdict = {}
    new = []

    def listfun(dict):
        print(dict)
        new.append(dict.copy())
        return new

    cursor = connection.cursor()
    locationwise = request.data.get("locationwise", None)
    locationname = request.data.get("locationname", None)
    groupby = request.data.get("groupby")
    where = request.data.get("where")
    previouslocation = request.data.get("previouslocation")
    clause = ""
    if (locationwise is not None) and (locationname == "all"):
        location = "ofc_" + locationwise
        clause = "WHERE " + previouslocation + "='" + where + "' "
        cursor.execute(
            f"""
   select {groupby} as location, count(distinct mr_id),count(prsnt_mtr_status='Ok' or NULL),count(rdng_ocr_status='Passed' or NULL),count(rdng_ocr_status='Failed' or NULL),count(prsnt_mtr_status='Meter Defective' or NULL),count(prsnt_mtr_status='Door Locked' or NULL),count(mr_id)
    from readingmaster {clause} and extract(month from reading_date_db)='{month}' and {groupby}!=''  GROUP BY {groupby}
    """
        )
    if (locationwise is not None) and (locationname != "all"):
        location = "ofc_" + locationwise
        clause = "WHERE " + locationwise + "='" + locationname + "' "

        cursor.execute(
            f"""
   select {groupby} as location, count(distinct mr_id),count(prsnt_mtr_status='Ok' or NULL),count(rdng_ocr_status='Passed' or NULL),count(rdng_ocr_status='Failed' or NULL),count(prsnt_mtr_status='Meter Defective' or NULL),count(prsnt_mtr_status='Door Locked' or NULL),count(mr_id)
    from readingmaster {clause} and extract(month from reading_date_db)='{month}' and {groupby}!='' GROUP BY {groupby}
    """
        )

    result = cursor.fetchall()
    print("result", result)
    try:
        for row in result:
            locationname = row[0]
            total = row[7]
            mrid = row[1]
            okreadings = row[2]
            OcrReadings = row[3]
            Ocrwithexception = row[4]
            meterDefective = row[5]
            doorLocked = row[6]
            okreadpercent = math.floor((okreadings / total) * 100)
            ocrreadingpercent = math.floor(
                ((OcrReadings / okreadings) if okreadings else 0) * 100
            )
            ocrwithexceppercent = math.floor(
                ((Ocrwithexception / okreadings) if okreadings else 0) * 100
            )
            meterdefectivepercent = math.floor((meterDefective / total) * 100)
            doorlockedpercent = math.floor((doorLocked / total) * 100)
            newdict["locationname"] = row[0]
            newdict["mrid"] = row[1]
            newdict["okreadings"] = row[2]
            newdict["okreadingspercent"] = okreadpercent
            newdict["OcrReadings"] = row[3]
            newdict["OcrReadingspercent"] = ocrreadingpercent
            newdict["Ocrwithexception"] = row[4]
            newdict["Ocrwithexceptionpercent"] = ocrwithexceppercent
            newdict["meterDefective"] = row[5]
            newdict["meterDefectivepercent"] = meterdefectivepercent
            newdict["doorLocked"] = row[6]
            newdict["doorLockedpercent"] = doorlockedpercent
            newdict["total"] = total
            data = listfun(newdict)
        return Response(data)
    except:
        return Response([])


def to_geojson(entries):
    features = []
    for entry in entries:
        if entry["geo_lat"] != "" and entry["geo_long"] != "":
            geolat = entry["geo_lat"] = float(entry["geo_lat"])
            geolong = entry["geo_long"] = float(entry["geo_long"])
            point = Point([entry["geo_long"], entry["geo_lat"]])

            del entry["geo_lat"]
            del entry["geo_long"]
            feature = Feature(geometry=point, properties=entry)
            features.append(feature)
    crs = {"type": "name", "properties": {
        "name": "urn:ogc:def:crs:OGC:1.3:CRS84"}}
    feature_collection = FeatureCollection(crs=crs, features=features)
    return feature_collection


def filtermethod():
    pass


@api_view(["GET"])
def tester(request):
    pass

@api_view(["GET"])
def clusterstest(request):
    mrid = request.query_params.get("mrid")
    today = date.today()
    clause = ""
    if mrid is not None:
        clause = "and mr_id" "='" + mrid + "' "
        print(clause)

    cursor = connection.cursor()
    query = f"""
    select mr_id,rdng_date,cons_name,geo_lat,geo_long,prsnt_mtr_status,rdng_ocr_status,rdng_img from readingmaster where reading_date_db='{today}' {clause}
    """
    cursor.execute(query)
    result = dictfetchall(cursor)
    return Response(result)

@api_view(["GET"])
def get_meter_summarytest(request):
    filters = request.query_params.dict()
    print("filters", filters)
    new = []

    def listfun(dict):
        # print(dict)
        new.append(dict.copy())
        return new

    newdict = {}
    cursor = connection.cursor()
    clause = f""""""

    if filters:
        clause += "WHERE "
        for i, (key, value) in enumerate(filters.items()):
            if i > 0:
                clause += "AND "
            if key == "month":
                print("key['month']", key)
                print("value['month']", value)
                month = value.split("-")[1]
                key = "extract(month from reading_date_db)"
                value = month

            clause += f"{key}='{value}'"
        print("clause", clause)
    # print("query",query)
    query = f"""select readingmaster.mr_id,count(readingmaster.id),count(readingmaster.prsnt_mtr_status='Ok' or NULL),count(readingmaster.rdng_ocr_status='Passed' or NULL),count(readingmaster.rdng_ocr_status='Failed' or NULL),count(readingmaster.prsnt_mtr_status='Meter Defective' or NULL),count(readingmaster.prsnt_mtr_status='Door Locked' or NULL)
            from readingmaster {clause}  group by readingmaster.mr_id
        """
    print("query", query)
    cursor.execute(query)
    results = cursor.fetchall()
    try:
        for row in results:
            total = row[1]
            okreadings = row[2]
            ocrreadings = row[3]
            ocrwithexcep = row[4]
            meterdefective = row[5]
            doorlocked = row[6]

            # Percentage
            okreadpercent = math.floor(((okreadings / total) * 100))
            ocrreadingpercent = math.floor(
                (((ocrreadings / okreadings) if okreadings else 0) * 100)
            )
            ocrwithexceppercent = math.floor(
                (((ocrwithexcep / okreadings) if okreadings else 0) * 100)
            )
            meterdefectivepercent = math.floor(
                ((meterdefective / total) * 100))
            doorlockedpercent = math.floor(((doorlocked / total) * 100))

            # add to dictionary
            newdict["mrid"] = row[0]
            newdict["totalReadings"] = row[1]
            newdict["OKreadings"] = okreadings
            newdict["OKreadingspercent"] = okreadpercent
            newdict["OCRReadings"] = ocrreadings
            newdict["OCRReadingspercent"] = ocrreadingpercent
            newdict["OCRwithException"] = ocrwithexcep
            newdict["OCRwithExceptionpercent"] = ocrwithexceppercent
            newdict["MeterDefective"] = meterdefective
            newdict["MeterDefectivepercent"] = meterdefectivepercent
            newdict["DoorLocked"] = doorlocked
            newdict["DoorLockedpercent"] = doorlockedpercent
            # add to list
            newdata = listfun(newdict)
        return Response(newdata)
    except:
        return Response([])


@api_view(["GET"])
def consumerwisemap(request):
    consacno = request.query_params.get("consacno")
    clause = ""

    if consacno is not None:
        clause = "where cons_ac_no" "='" + consacno + "' "
        print(clause)

    cursor = connection.cursor()
    query = f"""
    select mr_id,rdng_date,cons_name,geo_lat,geo_long,prsnt_mtr_status,rdng_ocr_status,rdng_img from readingmaster  {clause}
    """
    cursor.execute(query)
    result = dictfetchall(cursor)
    return Response(result)


@api_view(["POST"])
def geoclusternew(request):
    print("qwerty")
    data = request.data.get("filters", None)
    today = date.today()
    clause = ""
    # try:
    if data:
        print("rtyui")
        clause += "WHERE"
        for i, (key, value) in enumerate(data.items()):
            if i > 0:
                clause += " AND "
            if key == "prsnt_mtr_status":
                clause += f" {key}='{value}'"
            if key == "bl_agnc_name":
                clause += f" {key}='{value}'"
    cursor = connection.cursor()
    if data:
        query = f"""
        select m.mr_id as id,m.geo_lat,m.geo_long from readingmaster m {clause} and reading_date_db='{today}'
    """
    else:
        query = f"""
        select m.mr_id as id,m.geo_lat,m.geo_long from readingmaster m where reading_date_db='{today}'
        """
    print(query)
    cursor.execute(query)
    result = dictfetchall(cursor)
    geo = to_geojson(result)
    return Response(geo)

@api_view(["POST"])
def meterWiseReportUpdate(request):
    data = request.data.get("filters", None)
    groupby = request.data.get("groupby", None)
    print(request.data)

    if data is None:
        return Response([])

    else:
        today = date.today()
        clause = ""
        clause += "WHERE "
        if (
            data.get("month", "") == ""
            and data.get("startdate", "") == ""
            and data.get("enddate", "") == ""
        ):
            clause += f"extract(month from m.reading_date_db) = '{date.today().month}' AND extract(year from m.reading_date_db) = '{date.today().year}' AND "
        for i, (key, value) in enumerate(data.items()):
            if key == "month" and value:
                year = value.split("-")[0]
                month = value.split("-")[1]
                clause += f"extract(month from m.reading_date_db) = '{month}' AND extract(year from m.reading_date_db) = '{year}' AND "

            elif key == "startdate" and value:
                clause += f"extract(day from m.reading_date_db) BETWEEN '{data['startdate']}' AND "

            elif key == "enddate" and value:
                clause += f"'{data['enddate']}' AND "

            elif (
                key
                in [
                    "ofc_discom",
                    "ofc_zone",
                    "ofc_circle",
                    "ofc_division",
                    "ofc_subdivision",
                    "bl_agnc_name",
                ]
                and value
            ):
                clause += f"m.{key} ='{data[key]}' AND "

        if clause[-4:-1] == "AND":
            clause = clause[0:-4]
    print(clause)

    try:
        cursor = connection.cursor()
        if groupby != "ofc_section":
            query = f"""
                    SELECT m.{groupby} as location, count(m.id) as id,
            count(m.prsnt_mtr_status='Ok' or null) as Ok,
            CASE WHEN count(m.prsnt_mtr_status='Ok' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_mtr_status='Ok' or null) as float)
                        / cast(count(m.id) as float) * 100)::numeric, 2) END as ok_persent,
            count(m.rdng_ocr_status='Passed' or null) as ocr_Passed,
            CASE WHEN count(m.rdng_ocr_status='Passed' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.rdng_ocr_status='Passed' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as ocr_Passed_persent,
            count(m.prsnt_rdng_ocr_excep='Parameters Incorrect' or null) as Parameters_Incorrect,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Parameters Incorrect' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Parameters Incorrect' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Parameters_Incorrect_Persent,
            count(m.prsnt_rdng_ocr_excep='Parameters Unclear' or null) as Parameters_Unclear,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Parameters Unclear' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Parameters Unclear' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Parameters_Unclear_Persent,
            count(m.prsnt_rdng_ocr_excep='Parameters Unavailable' or null) as Parameters_Unavailable,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Parameters Unavailable' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Parameters Unavailable' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Parameters_Unavailable_Persent,
            count(m.prsnt_rdng_ocr_excep='Image Invalid' or null) as Image_Invalid,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Image Invalid' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Image Invalid' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Image_Invalid_Persent,
            count(m.prsnt_rdng_ocr_excep='Image Unclear' or null) as Image_Unclear,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Image Unclear' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Image Unclear' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Image_Unclear_Persent,
            count(m.prsnt_rdng_ocr_excep='Image Spoofed' or null) as Image_Spoofed,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Image Spoofed' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Image Spoofed' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Image_Spoofed_Persent,
            count(m.prsnt_rdng_ocr_excep='Image Stain on Decimal' or null) as Image_Stain_OnDecimal,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Image Stain on Decimal' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Image Stain on Decimal' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Image_Stain_On_Decimal_persent,
            count(m.prsnt_rdng_ocr_excep='Meter Mismatched' or null) as Meter_Mismatched,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Meter Mismatched' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Meter Mismatched' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Meter_Mismatched_persent,
            count(m.prsnt_rdng_ocr_excep='Meter On Height' or null) as Meter_On_Height,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Meter On Height' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Meter On Height' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Meter_On_Height_Persent,
            count(m.prsnt_rdng_ocr_excep='Meter Dirty' or null) as Meter_Dirty,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Meter Dirty' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Meter Dirty' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Meter_Dirty_Persent,
            count(m.prsnt_rdng_ocr_excep='Meter Display Broken' or null) as Meter_Display_Broken,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Meter Display Broken' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Meter Display Broken' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Meter_Display_Broken_Persent,
            count(m.prsnt_rdng_ocr_excep='Daylight Reflection On Meter' or null) as Daylight_Reflection_On_Meter,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Daylight Reflection On Meter' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Daylight Reflection On Meter' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Daylight_Reflection_On_Meter_Persent,
            count(m.prsnt_rdng_ocr_excep='Backlight Reflection' or null) as Backlight_Reflection,
            CASE WHEN count(m.prsnt_rdng_ocr_excep='Backlight Reflection' or null) = 0 THEN 0
            ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Backlight Reflection' or null) as float)
                        / cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2) END as Backlight_Reflection_Persent
            from readingmaster m {clause} and m.{groupby}!='' group by m.{groupby}
                """
        else:
            query = f"""
                select
                    o.sectionname as location,
                    count(m.id) as id,
                    count(m.prsnt_mtr_status='Ok' or null) as Ok,
                    count(m.rdng_ocr_status='Passed' or null) as ocr_Passed,
                    count(m.prsnt_rdng_ocr_excep='Parameters Incorrect' or null) as Parameters_Incorrect,
                    count(m.prsnt_rdng_ocr_excep='Parameters Unclear' or null) as Parameters_Unclear,
                    count(m.prsnt_rdng_ocr_excep='Parameters Unavailable' or null) as Parameters_Unavailable,
                    count(m.prsnt_rdng_ocr_excep='Image Invalid' or null) as Image_Invalid,
                    count(m.prsnt_rdng_ocr_excep='Image Unclear' or null) as Image_Unclear,
                    count(m.prsnt_rdng_ocr_excep='Image Spoofed' or null) as Image_Spoofed,
                    count(m.prsnt_rdng_ocr_excep='Image Stain on Decimal' or null) as Image_Stain_OnDecimal,
                    count(m.prsnt_rdng_ocr_excep='Meter Mismatched' or null) as Meter_Mismatched,
                    count(m.prsnt_rdng_ocr_excep='Meter On Height' or null) as Meter_On_Height,
                    count(m.prsnt_rdng_ocr_excep='Meter Dirty' or null) as Meter_Dirty,
                    count(m.prsnt_rdng_ocr_excep='Meter Display Broken' or null) as Meter_Display_Broken,
                    count(m.prsnt_rdng_ocr_excep='Daylight Reflection On Meter' or null) as Daylight_Reflection_On_Meter,
                    count(m.prsnt_rdng_ocr_excep='Backlight Reflection' or null) as Backlight_Reflection,
                    CASE WHEN
                        count(m.prsnt_mtr_status='Ok' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_mtr_status='Ok' or null) as float)/cast(count(m.id) as float) * 100)::numeric, 2)
                    END as ok_persent,

                    CASE WHEN
                        count(m.rdng_ocr_status='Passed' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.rdng_ocr_status='Passed' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as ocr_Passed_persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Parameters Incorrect' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Parameters Incorrect' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Parameters_Incorrect_Persent,
                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Parameters Unclear' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Parameters Unclear' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Parameters_Unclear_Persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Parameters Unavailable' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Parameters Unavailable' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Parameters_Unavailable_Persent,
                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Image Invalid' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Image Invalid' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Image_Invalid_Persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Image Unclear' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Image Unclear' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Image_Unclear_Persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Image Spoofed' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Image Spoofed' or null) as float)/ cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Image_Spoofed_Persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Image Stain on Decimal' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Image Stain on Decimal' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Image_Stain_On_Decimal_persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Meter Mismatched' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Meter Mismatched' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Meter_Mismatched_persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Meter On Height' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Meter On Height' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Meter_On_Height_Persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Meter Dirty' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Meter Dirty' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Meter_Dirty_Persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Meter Display Broken' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Meter Display Broken' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Meter_Display_Broken_Persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Daylight Reflection On Meter' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Daylight Reflection On Meter' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Daylight_Reflection_On_Meter_Persent,

                    CASE WHEN
                        count(m.prsnt_rdng_ocr_excep='Backlight Reflection' or null) = 0 THEN 0 ELSE ROUND((cast(count(m.prsnt_rdng_ocr_excep='Backlight Reflection' or null) as float)/cast(count(m.prsnt_mtr_status='Ok' or null) as float) * 100)::numeric, 2)
                    END as Backlight_Reflection_Persent

                from readingmaster m
                join
                    (
                        SELECT DISTINCT sectioncode, MAX(sectionname) AS sectionname
                        FROM office
                        GROUP BY sectioncode
                    ) o ON m.ofc_section = o.sectioncode
                {clause} and o.sectionname!=''
                group by o.sectionname
                """
        # print(clause)
        print(query)
        cursor.execute(query)
        person_objects = dictfetchall(cursor)
        return Response(person_objects)
    except Exception as e:
        print(e)
        return Response([])


@api_view(["POST"])
def meterWiseReportconsumer(request):
    # pagesize = request.data.get("pagesize",)
    # page = (request.data.get("page",))
    # offset=(int(pagesize) * int(page))-int(pagesize)
    data = request.data.get("filters", None)

    clause = ""
    try:
        if data:
            clause += "WHERE "
            if data.get("month", "") == "":
                clause += f"extract(month from m.reading_date_db) = '{date.today().month}' AND extract(year from m.reading_date_db) = '{date.today().year}' AND "

            for i, (key, value) in enumerate(data.items()):
                if key == "month" and value:
                    year = value.split("-")[0]
                    month = value.split("-")[1]
                    clause += f"extract(month from m.reading_date_db) = '{month}' AND extract(year from m.reading_date_db) = '{year}' AND "

                elif key == "startdate" and value:
                    clause += f"m.reading_date_db BETWEEN '{data['startdate']}' AND "

                elif key == "enddate" and value:
                    clause += f"'{data['enddate']}' AND "

                elif (
                    key
                    in [
                        "ofc_discom",
                        "ofc_zone",
                        "ofc_circle",
                        "ofc_division",
                        "ofc_subdivision",
                        "bl_agnc_name",
                    ]
                    and value
                ):
                    clause += f"m.{key} ='{data[key]}' AND "

            if clause[-4:-1] == "AND":
                clause = clause[0:-4]

                # clause += f" {key}='{value}'"
            cursor = connection.cursor()
            #
            query = f"""
                Select distinct m.cons_ac_no as consAccountNumber, m.cons_name as consName,m.ofc_discom as discom,m.ofc_zone as zone,
                m.ofc_circle as circle,m.ofc_division as division,m.ofc_subdivision as subdivision,o.sectionname as section,m.rdng_img as readingImg,
                m.md_img as mdImg from readingmaster m join office o ON m.ofc_section=o.sectioncode {clause}
                """

            print(clause)
            print(query)
            # serializer=meterWiseReportconsumerSerializer(query,many=True)
            cursor.execute(query)
            person_objects = dictfetchall(cursor)
            return Response({"results": person_objects})
            # print(serializer.data)
            # return Response({"results":serializer.data})
    except:
        return Response([])


# ------------------------------------------------------ NEW APIS WITH MULTIPLE FILTERS-------------------------------------------------------------------------------------------

@api_view(["POST"])
def supervisorlogin(request):
    number = request.data.get("supervisor_number")
    password = request.data.get("password")

    if not number or not password:
        return Response({
            "status": False,
            "message": "supervisor_number and password required",
            "accessToken": "",
            "user": None
        })

    user = SupervisorLogin.objects.filter(
        supervisor_number=number,
        password=password
    ).first()   # <-- FIX

    if not user:
        return Response({
            "status": False,
            "message": "login failed",
            "accessToken": "",
            "user": None
        })

    return Response({
        "status": True,
        "message": "login successful",
        "accessToken": "",
        "user": {
            "supervise_name": user.supervisor_name or "",
            "supervise_number": user.supervisor_number or "",
            "is_admin": user.is_admin,
            "designation": "Supervisor",
            "division": user.ofc_division or "",
            "subdivison": user.ofc_subdivision or ""
        }
    })

#indra
from django.db import connection
from rest_framework.decorators import api_view
from rest_framework.response import Response
from rest_framework import status


@api_view(["POST"])
def get_meter_reader_subdivision(request):

    mr_id = request.data.get("mr_id")

    # mr_id not provided
    if mr_id is None or mr_id == "":
        return Response(
            {
                "success": False,
                "message": "mr_id is required",
                "ofc_subdivision": None
            },
            status=status.HTTP_400_BAD_REQUEST
        )

    # Validate mr_id before sending it to PostgreSQL
    try:
        mr_id = int(mr_id)
    except (ValueError, TypeError):
        return Response(
            {
                "success": True,
                "message": "Meter reader mapping not found",
                "mr_id": request.data.get("mr_id"),
                "ofc_subdivision": None
            },
            status=status.HTTP_200_OK
        )

    query = """
        SELECT ofc_subdivision
        FROM public.meter_reader_mapping
        WHERE meter_reader_id = %s
        LIMIT 1
    """

    with connection.cursor() as cursor:
        cursor.execute(query, [mr_id])
        row = cursor.fetchone()

    # Numeric ID but doesn't exist in DB
    if not row:
        return Response(
            {
                "success": True,
                "message": "Meter reader mapping not found",
                "mr_id": mr_id,
                "ofc_subdivision": None
            },
            status=status.HTTP_200_OK
        )

    return Response(
        {
            "success": True,
            "message": "Meter reader mapping found",
            "mr_id": mr_id,
            "ofc_subdivision": row[0]
        },
        status=status.HTTP_200_OK
    )



# deeepak
from datetime import datetime
from django.db import connection
from rest_framework.decorators import api_view
from rest_framework.response import Response
 
@api_view(["POST"])
def supervisorlocation(request):
    supervisor_number = request.data.get("supervisor_number")
    geo_lat = request.data.get("geo_lat")
    geo_long = request.data.get("geo_long")
    date_str = request.dataget = request.data.get("date")
 
    if not all([supervisor_number, geo_lat, geo_long, date_str]):
        return Response({"status": False, "message": "Missing fields"}, status=400)
 
    # Parse datetime
    try:
        dt = datetime.strptime(date_str, "%Y-%m-%d %I:%M:%S %p")
        date = dt.date()
        time = dt.strftime("%H:%M:%S")
 
    except:
        return Response({"status": False, "message": "Invalid date format"}, status=400)
 
    try:
        with connection.cursor() as cursor:
            cursor.execute("""
                INSERT INTO supervsiorlocation
                    (supervisor_number, date, meta, created_at, updated_at)
                VALUES (
                    %s,%s,
                    jsonb_build_object(
                        'path', jsonb_build_array(
                            jsonb_build_object(
                                'time', %s,
                                'lat', %s,
                                'lng', %s
                            )
                        ),
                        'last_seen', %s,
                        'total_points', 1
                    ),
                    NOW(),NOW()
                )
 
                ON CONFLICT (supervisor_number, date)
                DO UPDATE SET
                    meta =
                        jsonb_set(
                            jsonb_set(
                                jsonb_set(
                                    supervsiorlocation.meta,
                                    '{path}',
                                    (supervsiorlocation.meta->'path') || jsonb_build_array(
                                        jsonb_build_object(
                                            'time', %s,
                                            'lat', %s,
                                            'lng', %s
                                        )
                                    )
                                ),
                                '{last_seen}',
                                to_jsonb(%s::text)
                            ),
                            '{total_points}',
                            to_jsonb((supervsiorlocation.meta->>'total_points')::int + 1)
                        ),
                    updated_at = NOW();
            """, [
                supervisor_number, date,
                time, geo_lat, geo_long, time,
                time, geo_lat, geo_long, time
            ])
        return Response({"status": True, "message": "Location recorded"})
    except Exception as e:
        print("DB ERROR:", e)
        return Response({"status": False, "message": str(e)}, status=500)


@api_view(["POST"])
def gitnewmvcheck(request):
    pagesize = request.data.get("pagesize", None)
    page = request.data.get("page", 1)
    orderby = request.data.get("orderby", "DESC")
    filters = request.data.get("filters", {})
    print("filters...", filters)
    export_all = request.data.get("export_all", False)  # NEW FLAG

    offset = (int(pagesize) * int(page)) - int(pagesize) if pagesize else 0

    # Build filter clause
    clause_parts = []
    for key, value in filters.items():
        if key == "month":
            year, month = value.split("-")
            clause_parts.append(
                f"EXTRACT(month from m.reading_date_db) = '{month}'")
            clause_parts.append(
                f"EXTRACT(year from m.reading_date_db) = '{year}'")
        elif key == "startdate":
            clause_parts.append(
                f"EXTRACT(day from m.reading_date_db) >= '{value}'")
        elif key == "enddate":
            clause_parts.append(
                f"EXTRACT(day from m.reading_date_db) <= '{value}'")
        elif key == "mr_id":
            clause_parts.append(f"m.mr_id = '{value}'")
        elif key == "prsnt_mtr_status":
            clause_parts.append(f"m.prsnt_mtr_status = '{value}'")
        elif key == "reading_parameter_type":
            clause_parts.append(f"m.reading_parameter_type = '{value}'")
            clause_parts.append("m.rdng_ocr_status = 'Failed'")
        elif key == "searchdata":
            clause_parts.append(
                f"(m.mr_id = '{value}' OR m.cons_ac_no = '{value}' OR m.cons_name = '{value}')"
            )
        elif key == "rdng_ocr_status":
            if value == "OCR without Exception":
                clause_parts.append("m.rdng_ocr_status = 'Passed'")
            elif value == "OCR with Exception":
                exception_detail = filters.get("prsnt_rdng_ocr_excep")
                clause_parts.append("m.rdng_ocr_status = 'Failed'")
                if exception_detail:
                    clause_parts.append(
                        f"m.prsnt_rdng_ocr_excep = '{exception_detail}'")
        elif key == "bl_agnc_name":
            clause_parts.append(f"bl_agnc_name = '{value}'")
        elif key == "ofc_discom":
            clause_parts.append(f"ofc_discom = '{value}'")

    clause = " AND ".join(clause_parts)
    clause = f" AND {clause}" if clause else ""

    tablename = "readingmaster"  # Adjust if needed

    # Base SELECT
    query = f"""
        SELECT m.con_mtr_sl_no, m.mr_id as "mrId", m.rdng_date, m.prsnt_mtr_status, m.prsnt_ocr_rdng,
               m.prsnt_rdng, m.ocr_pf_status, pf_image, pf_manual_reading,
               m.cons_name, m.cons_ac_no, m.prsnt_md_rdng_ocr, m.rdng_ocr_status,
               m.rdng_img, m.prsnt_md_rdng, m.id, r."mrPhoto",
               m.prsnt_rdng_ocr_excep, m.reading_parameter_type
        FROM {tablename} m
        LEFT JOIN meterreaderregistration r on m.mr_id=r."mrId"
        WHERE (m.rdng_ocr_status_changed_by IS NULL OR m.rdng_ocr_status_changed_by=''
               OR m.rdng_ocr_status_changed_by ILIKE '%vapp%')
        AND m.rdng_img != '' {clause}
        ORDER BY m.rdng_date {orderby}
    """

    # Only apply LIMIT/OFFSET when NOT exporting all
    if not export_all and pagesize:
        query += f" LIMIT {pagesize} OFFSET {offset}"

    cursor = connection.cursor()
    cursor.execute(query)
    results = dictfetchall(cursor)

    if export_all:
        # No need to run count, just return all rows
        return Response({"count": len(results), "results": results})
    else:
        # Normal pagination → get total count
        query_total = f"""
            SELECT COUNT(*) FROM {tablename} m
            LEFT JOIN meterreaderregistration r on m.mr_id=r."mrId"
            WHERE (m.rdng_ocr_status_changed_by IS NULL OR m.rdng_ocr_status_changed_by=''
                   OR m.rdng_ocr_status_changed_by ILIKE '%vapp%')
            AND m.rdng_img != '' {clause}
        """
        cursor.execute(query_total)
        total_count = dictfetchall(cursor)[0]["count"]

        return Response({"count": total_count, "results": results})


@api_view(["POST"])
def clusterstestnew(request):
    data = request.data.get("filters", None)
    today = date.today()
    clause = ""
    # try:
    if data:
        # this below code is for supervisor location
        if "mr_id" in data:
            mr_id_value = data["mr_id"]
            if mr_id_value.startswith('SUP_'):
                supervisor_number = mr_id_value[4:]
                print("-------->>>>", today, supervisor_number)
                try:
                    with connection.cursor() as cursor:
                        cursor.execute("""
                            SELECT
                            jsonb_agg(
                                jsonb_build_object(
                                'geo_lat',  p->>'lat',
                                'geo_long', p->>'lng',
                                'time',     p->>'time'
                                )
                            ) AS path,
                            sl.supervisor_number,
                            sl.date
                            FROM supervsiorlocation sl
                            CROSS JOIN LATERAL jsonb_array_elements(sl.meta->'path') AS p
                            WHERE sl.supervisor_number = %s
                            AND sl.date = %s
                            GROUP BY sl.supervisor_number, sl.date
                        """, [supervisor_number, today])

                        row = cursor.fetchone()

                    if not row:
                        return Response([])

                    response_data = {
                        "supervisor_number": row[1],
                        "date": row[2],
                        "path": json.loads(row[0]) if isinstance(row[0], str) else row[0]
                    }

                    supervisor_login_data = SupervisorLogin.objects.filter(
                        supervisor_number=supervisor_number
                    ).values(
                        'supervisor_name', 'ofc_division', 'ofc_subdivision'
                    ).first()

                    if supervisor_login_data:
                        response_data.update(supervisor_login_data)
                    return Response(response_data)
                except Exception as e:
                    return Response({"error": str(e)}, status=500)

        clause += "WHERE "
        for i, (key, value) in enumerate(data.items()):
            if i > 0:
                clause += " AND "
            if key == "mr_id":
                clause += f" {key}='{value}'"
            if key == "bl_agnc_name":
                clause += f" {key}='{value}'"

        cursor = connection.cursor()
        query = f"""
        select mr_id,rdng_date,cons_name,geo_lat,geo_long,prsnt_mtr_status,rdng_ocr_status,prsnt_ocr_rdng,ocr_pf_reading,cons_ac_no,prsnt_md_rdng_ocr,prsnt_md_rdng,prsnt_rdng,qc_req,
        rdng_img from readingmaster {clause} AND reading_date_db='{today}'
        """
        print(query)
        cursor.execute(query)
        result = dictfetchall(cursor)
        return Response(result)


from datetime import date, datetime, timedelta
from django.db import connection
# from rest_framework.decorators import api_view
from rest_framework.response import Response

@api_view(["POST"])
def androidclusterstestnew(request):
    filters = request.data.get("filters", {})
    today = date.today()

    where_clauses = ["reading_date_db = %s"]
    params = [str(today)]

    # Dynamic filters
    if "mr_id" in filters:
        where_clauses.append("mr_id = %s")
        params.append(filters["mr_id"])

    if "bl_agnc_name" in filters:
        where_clauses.append("bl_agnc_name = %s")
        params.append(filters["bl_agnc_name"])

    where_sql = "WHERE " + " AND ".join(where_clauses)

    query = f"""
        SELECT DISTINCT ON (mr_id)
            mr_id, rdng_date, cons_name, geo_lat, geo_long,
            prsnt_mtr_status, rdng_ocr_status, rdng_img
        FROM readingmaster
        {where_sql}
        ORDER BY mr_id,
                 (geo_lat IS NULL OR geo_long IS NULL),  -- Prefer NOT NULL
                 rdng_date DESC                          -- Latest record
    """

    cursor = connection.cursor()
    cursor.execute(query, params)
    print("query:>", query)
    result = dictfetchall(cursor)

    return Response(result)

@api_view(["GET"])
def downloadexcel(request):

    supervisor_number = request.query_params.get("supervisor_number")

    # Fetch supervisor rows
    qs = SupervisorLogin.objects.filter(supervisor_number=supervisor_number)

    if not qs.exists():
        return Response({"status": False, "message": "Supervisor not found"})

    # All MR IDs
    mr_ids = list(qs.values_list("mr_id", flat=True))

    # Convert → 'MRC1','MRC2',...
    mr_sql_list = ",".join(f"'{x}'" for x in mr_ids)

    # Supervisor info
    sup = qs.first()
    ofc_division = sup.ofc_division
    ofc_subdivision = sup.ofc_subdivision

    datewise = request.query_params.get("datewise")
    today = datetime.now().date()
    month = datetime.now().month

    if datewise == "date1":
        clause = f"reading_date_db = '{today}'"
    else:
        clause = f"extract(Month from reading_date_db) = '{month}'"

    cursor = connection.cursor()

    # ----- SUMMARY -----
    query = f"""
        SELECT 
            COUNT(r.id) AS total_readings,
            COUNT(r.qc_req='Yes' OR NULL) AS qc_remaining,
            COUNT(r.qc_req='No' OR NULL) AS qc_done
        FROM readingmaster r
        WHERE 
            {clause}
            AND r.mr_id IN ({mr_sql_list})
    """

    print("query:", query)
    cursor.execute(query)
    summary = dictfetchall(cursor)

    # ----- MR WISE DATA -----
    query2 = f"""
        SELECT 
            DISTINCT(mr_id),
            CASE WHEN COUNT(prsnt_mtr_status='Ok' OR NULL) = 0 THEN 0
            ELSE ROUND(
                (
                    CAST(COUNT(rdng_ocr_status='Passed' OR NULL) AS FLOAT) /
                    CAST(COUNT(prsnt_mtr_status='Ok' OR NULL) AS FLOAT) * 100
                )::numeric
            , 2) END AS passed_percent,

            ROUND(
                (
                    CAST(COUNT(prsnt_mtr_status='Meter Defective' OR NULL) AS FLOAT) /
                    COALESCE(CAST(COUNT(mr_id) AS FLOAT),1) * 100
                )::numeric
            , 2) AS meter_defective_percent,

            ROUND(
                (
                    CAST(COUNT(prsnt_mtr_status='Door Locked' OR NULL) AS FLOAT) /
                    COALESCE(CAST(COUNT(mr_id) AS FLOAT),1) * 100
                )::numeric
            , 2) AS door_locked_percent,

            COUNT(id) AS mr_total_readings,
            COUNT(qc_req='Yes' OR NULL) AS mr_qc_remaining,
            COUNT(qc_req='No' OR NULL) AS mr_qc_done

        FROM readingmaster
        WHERE 
            {clause}
            AND mr_id IN ({mr_sql_list})
        GROUP BY mr_id
    """

    print("query2:", query2)
    cursor.execute(query2)
    rows = dictfetchall(cursor)

    # ---------- EXCEL ----------
    wb = Workbook()
    ws = wb.active

    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for item in rows:
            ws.append(list(item.values()))

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=mydata.xlsx"
    wb.save(response)
    return response

# indra
@api_view(["POST"])
def meterreaderDetails(request):
    pagesize = request.data.get("pagesize")
    page = request.data.get("page")
    offset = (int(pagesize) * int(page)) - int(pagesize)

    import time

    start = time.time()

    data = request.data.get("filters", None)
    clause = ""
    try:
        if data:
            clause += " WHERE "
            conditions = []
            for key, value in data.items():
                if key == "month":
                    year = value.split("-")[0]
                    month = value.split("-")[1]
                    conditions.append(
                        f"extract(month from reading_date_db) = '{month}' AND extract(year from reading_date_db) = '{year}'")

                if key == "startdate":
                    conditions.append(
                        f"extract(day from reading_date_db) BETWEEN '{data['startdate']}'")

                if key == "enddate":
                    conditions.append(f"'{data['enddate']}'")

                if key == "mr_id":
                    conditions.append(f"mr_id='{data['mr_id']}'")

                if key == "searchdata":
                    conditions.append(
                        f"(mr_id='{data['searchdata']}' OR cons_ac_no='{data['searchdata']}' OR cons_name='{data['searchdata']}')")

                if key == "rdng_ocr_status":
                    conditions.append(
                        f"rdng_ocr_status='{data['rdng_ocr_status']}'")
                if key == "Exception":
                    conditions.append(f"rdng_ocr_status='{value}'")

                if key == "prsnt_rdng_ocr_excep":
                    # CASE 1: Passed → get only passed rows
                    if value == "Passed":
                        conditions.append("rdng_ocr_status = 'Passed'")
                    # CASE 2: Failed (All)
                    elif value == "__FAILED__":
                        # Get rows where there IS an exception (not empty, not null)
                        conditions.append(
                            "TRIM(COALESCE(prsnt_rdng_ocr_excep, '')) <> ''")
                    # CASE 3: Failed + Specific Reason
                    else:
                        conditions.append(f"prsnt_rdng_ocr_excep = '{value}'")

                if key == "con_trf_cat":
                    conditions.append(f"con_trf_cat='{value}'")

                if key == 'prsnt_mtr_status':
                    conditions.append(f"prsnt_mtr_status='{value}'")

                if key == "bl_agnc_name":
                    conditions.append(f"bl_agnc_name='{data['bl_agnc_name']}'")

                if key == "Discom":
                    conditions.append(f"ofc_discom='{data['Discom']}'")
                # if key == "ofc_discom":
                #     conditions.append(f"ofc_discom='{data['ofc_discom']}'")

                if key == "zone":
                    conditions.append(f"ofc_zone='{data['zone']}'")

                if key == "circle":
                    conditions.append(f"ofc_circle='{data['circle']}'")

                if key == "Division":
                    conditions.append(f"ofc_division='{data['Division']}'")

                if key == "Subdivision":
                    conditions.append(
                        f"ofc_subdivision='{data['Subdivision']}'")

                if key == "Section":
                    conditions.append(f"ofc_section='{data['Section']}'")

            # Join all conditions using 'AND'
            clause += " AND ".join(conditions)

            selected_month = data.get("month", None)
            today = datetime.now()
            this_month = today.strftime("%Y-%m")
            previous_month = (
                today - timedelta(days=today.day)).strftime("%Y-%m")
            tablename = "prevmonthsdata" if selected_month not in {
                this_month, previous_month} else "readingmaster"

            cursor = connection.cursor()
            query = f"""
                SELECT mr_id, cons_ac_no, bl_agnc_name, abnormality, cons_name, con_trf_cat, con_mtr_sl_no,
                mr_rmrk, prsnt_mtr_status, prsnt_rdng, prev_rdng, prsnt_md_rdng, prev_md, ocr_pf_reading,
                prev_pf_rdng, rdng_date, prev_rdng_date, rdng_img, md_img, rdng_ocr_status,
                CASE
                    WHEN rdng_ocr_status = 'Passed' THEN 'Passed'
                    ELSE COALESCE(NULLIF(TRIM(prsnt_rdng_ocr_excep), ''), '')
                END AS prsnt_rdng_ocr_excep,
                md_ocr_excep, qc_req FROM {tablename} {clause} ORDER BY rdng_date DESC LIMIT {pagesize} OFFSET {offset}
                """
            print("QUERY!", query)
            cursor.execute(query)
            person_objects = dictfetchall(cursor)

            query2 = f"""
                SELECT COUNT(*) AS total_count FROM {tablename} {clause}
                """
            print("QUERY!", query2)

            cursor.execute(query2)
            rows = cursor.fetchone()

            print(time.time() - start)
            return Response({"result": person_objects, "count": rows[0]})
        else:
            # No filters present, return empty response
            return Response({"result": [], "count": 0})

    except Exception as e:
        print(e)  # Log the error for debugging purposes
        return Response({"result": [], "count": 5})


@api_view(["POST"])
def cons_wise_details_with_search(request):
    new = []

    def listfun(dict):
        new.append(dict.copy())
        return new

    newdict = {}
    clause = ""
    cursor = connection.cursor()
    filters = request.data.get("filters", None)
    try:
        if filters:
            clause += "WHERE "
            for i, (key, value) in enumerate(filters.items()):
                if i > 0:
                    clause += "AND "
                if key == "month":
                    year = value.split("-")[0]
                    month = value.split("-")[1]
                    print("month", month)
                    clause += f"extract(month from reading_date_db) = '{month}' AND extract(year from reading_date_db) = '{year}'"
                if key == "startdate":
                    clause += f"reading_date_db BETWEEN '{filters['startdate']}'"

                if key == "enddate":
                    clause += f"'{filters['enddate']}'"
                if key == "cons_ac_no":
                    clause += f"cons_ac_no='{value}'"

    except:
        pass
    if clause != "":
        query = f"""SELECT  m.ofc_discom,m.ofc_zone,m.ofc_circle,m.ofc_division,m.ofc_subdivision,m.ofc_section,
    m.id,m.cons_name,m.cons_ac_no,cons_address,m.cons_ph_no,m.con_trf_cat,m.mr_unit,
    r."mrId",r."mrName",r."mrPhone",r."mrPhoto" as avatar,m.con_mtr_sl_no,
    m.rdng_date,m.prsnt_mtr_status,m.prsnt_md_rdng,m.ocr_pf_reading,m.abnormality,m.prsnt_rdng_ocr_excep,m.md_ocr_excep,m.mr_rmrk,m.qc_req,m.ai_mdl_ver,m.ph_name,m.cmra_res,m.andr_ver,m.reading_date_db,
    m.rdng_img,m.md_img,m.pf_image,m.prsnt_ocr_rdng,m.prsnt_rdng,m.prsnt_md_rdng_ocr,m.rdng_ocr_status,m.kvah_rdng,m.kvah_img 
      FROM
    readingmaster m left outer join meterreaderregistration r on m.mr_id=r."mrId" {clause}"""
    else:
        return Response({"MSG": "PROVIDE CONSUMER ACOUNT NUMBER"})

    print("QUERY------>", query)
    cursor.execute(query)
    results = dictfetchall(cursor)
    print(results)
    return Response(results)

@api_view(["POST"])
def cons_passed(request):
    cons_ac_no = request.data["cons_ac_no"]
    mr_id = request.data["mrId"]

    cursor = connection.cursor()
    query = f"""select distinct(rdng_ocr_status) from readingmaster where rdng_ocr_status='Passed'  and cons_ac_no='{cons_ac_no}' and mr_id='{mr_id}' and manual_update_flag isnull and  qc_done !='byLambda'
    """

    ocrstatus = ""
    cursor.execute(query)
    result = cursor.fetchall()
    if len(result) > 0:
        if result[0][0] == "Passed":
            ocrstatus = "Passed"

    else:
        ocrstatus = "THERE IS NO OCR PASSED FOR THIS CONSUMER"

    return Response({"status": ocrstatus})


#indra

from dateutil.relativedelta import relativedelta


# ----------------------------------------TruereadQCAPP--------------------------------

@api_view(['POST'])
def divisiondata(request):
    zone = request.data.get("zone", None)
    month = request.data.get("month", None)

    try:
        if month:
            start_date = datetime.strptime(month, "%Y-%m").replace(day=1)
            if start_date.month == 12:
                end_date = start_date.replace(
                    month=1, year=start_date.year + 1) - timedelta(days=1)
            else:
                end_date = start_date.replace(
                    month=start_date.month + 1) - timedelta(days=1)
        else:
            return Response({"error": "Month is required in 'YYYY-MM' format."}, status=400)
    except ValueError:
        return Response({"error": "Invalid month format. Expected 'YYYY-MM'."}, status=400)

    query = "SELECT DISTINCT ofc_division FROM readingmaster WHERE reading_date_db BETWEEN %s AND %s"
    params = [start_date, end_date]

    if zone:
        query += " AND ofc_zone = %s"
        params.append(zone)

    with connection.cursor() as cursor:
        cursor.execute(query, params)
        divisions = [row[0] for row in cursor.fetchall()]

    return Response(divisions if divisions else [])

@api_view(['POST'])
def search_by_mr(request):
    month = request.data.get('month', None)
    mrid = request.data.get('mrid', None)
    now = datetime.now()
    clause = ''
    where_clauses = []

    if month:
        month = month.split('-')[1]
        where_clauses.append(f"EXTRACT(MONTH FROM reading_date_db)='{month}'")
    else:
        current_month = now.month
        where_clauses.append(
            f"EXTRACT(MONTH FROM m.reading_date_db) = '{current_month}'")
    if mrid:
        where_clauses.append(f"m.mr_id like '{mrid}%'")
    if where_clauses:
        clause = ' AND '.join(where_clauses)
    cursor = connection.cursor()
    query = (
        f'''select distinct mr_id from readingmaster m left outer join meterreaderregistration r on m.mr_id=r."mrId" where {clause}''')
    print("query", query)
    cursor.execute(query)
    person_objects = dictfetchall(cursor)

    if not person_objects:
        return Response({
            "status": False,
            "message": "MR not found",
            "results": person_objects,
        })

    return Response({
        "status": True,
        "message": "MR Fetched Successfully",
        "results": person_objects,
    })


@api_view(['POST'])
def userdashboard(request):
    user = request.data.get('user', None)
    now = datetime.now()
    today = date.today()
    current_week = today.isocalendar()[1]
    current_month = now.month
    last_month = (now.month - 1) if now.month > 1 else 12

    query_today = f'''
    SELECT
            period,
            SUM(CASE WHEN rdng_ocr_status = 'Passed' THEN 1 ELSE 0 END) AS yes_count,
            SUM(CASE WHEN rdng_ocr_status = 'Failed' AND prsnt_rdng_ocr_excep != 'Spoofed Image' THEN 1 ELSE 0 END) AS no_count,
            SUM(CASE WHEN rdng_ocr_status = 'Failed' AND prsnt_rdng_ocr_excep = 'Spoofed Image' THEN 1 ELSE 0 END) AS spoof_count,
            SUM(CASE WHEN rdng_ocr_status = 'Failed' AND qc_rmrk = 'MR Fault' THEN 1 ELSE 0 END) AS mrfault_count,
            COUNT(*) AS tot_count
        FROM (
            SELECT
                qc_rmrk,
                rdng_ocr_status,
                prsnt_rdng_ocr_excep,
                TO_DATE(date_qc, 'YYYY-MM-DD') AS date_qc,
                CASE
                    WHEN TO_DATE(date_qc, 'YYYY-MM-DD') = %s THEN 'today'
                    ELSE 'other'
                END AS period
            FROM readingmaster
            WHERE mtr_excep_img = %s
        ) AS subquery
        WHERE period != 'other'
        GROUP BY period
    '''

    query_week = f'''
    SELECT
            period,
            SUM(CASE WHEN rdng_ocr_status = 'Passed' THEN 1 ELSE 0 END) AS yes_count,
            SUM(CASE WHEN rdng_ocr_status = 'Failed' AND prsnt_rdng_ocr_excep != 'Spoofed Image' THEN 1 ELSE 0 END) AS no_count,
            SUM(CASE WHEN rdng_ocr_status = 'Failed' AND prsnt_rdng_ocr_excep = 'Spoofed Image' THEN 1 ELSE 0 END) AS spoof_count,
            SUM(CASE WHEN rdng_ocr_status = 'Failed' AND qc_rmrk = 'MR Fault' THEN 1 ELSE 0 END) AS mrfault_count,
            COUNT(*) AS tot_count
        FROM (
            SELECT
                qc_rmrk,
                rdng_ocr_status,
                prsnt_rdng_ocr_excep,
                TO_DATE(date_qc, 'YYYY-MM-DD') AS date_qc,
                CASE
                    WHEN EXTRACT(WEEK FROM TO_DATE(date_qc, 'YYYY-MM-DD')) = %s THEN 'week'
                    ELSE 'other'
                END AS period
            FROM readingmaster
            WHERE mtr_excep_img = %s
        ) AS subquery
        WHERE period != 'other'
        GROUP BY period
    '''

    query = f'''
        SELECT
            period,
            SUM(CASE WHEN rdng_ocr_status = 'Passed' THEN 1 ELSE 0 END) AS yes_count,
            SUM(CASE WHEN rdng_ocr_status = 'Failed' AND prsnt_rdng_ocr_excep != 'Spoofed Image' THEN 1 ELSE 0 END) AS no_count,
            SUM(CASE WHEN rdng_ocr_status = 'Failed' AND prsnt_rdng_ocr_excep = 'Spoofed Image' THEN 1 ELSE 0 END) AS spoof_count,
            SUM(CASE WHEN rdng_ocr_status = 'Failed' AND qc_rmrk = 'MR Fault' THEN 1 ELSE 0 END) AS mrfault_count,
            COUNT(*) AS tot_count
        FROM (
            SELECT
                qc_rmrk,
                rdng_ocr_status,
                prsnt_rdng_ocr_excep,
                TO_DATE(date_qc, 'YYYY-MM-DD') AS date_qc,
                CASE
                    WHEN extract(month from TO_DATE(date_qc, 'YYYY-MM-DD')) = %s THEN 'month'
                    WHEN extract(month from TO_DATE(date_qc, 'YYYY-MM-DD')) = %s THEN 'lastmonth'
                    ELSE 'other'
                END AS period
            FROM readingmaster
            WHERE mtr_excep_img = %s
        ) AS subquery
        WHERE period != 'other'
        GROUP BY period
    '''

    print("Query:", query_today, query)
    print("Parameters:", [today, current_week,
          current_month, last_month, f'vapp_{user}'])

    with connection.cursor() as cursor:

        cursor.execute(query_today, [today, f'vapp_{user}'])
        result_today = cursor.fetchone()

        cursor.execute(query_week, [current_week, f'vapp_{user}'])
        result_week = cursor.fetchone()

        cursor.execute(
            query, [current_month, last_month, f'vapp_{user}'])
        result_rows = cursor.fetchall()

    response_data = {
        "today": {"tot_count": 0, "yes_per": 0, "no_per": 0, "spoof_per": 0, "mrfault_per": 0},
        "week": {"tot_count": 0, "yes_per": 0, "no_per": 0, "spoof_per": 0, "mrfault_per": 0},
        "month": {"tot_count": 0, "yes_per": 0, "no_per": 0, "spoof_per": 0, "mrfault_per": 0},
        "lastmonth": {"tot_count": 0, "yes_per": 0, "no_per": 0, "spoof_per": 0, "mrfault_per": 0},
    }

    print(result_rows)

    if result_today:
        period = result_today[0]
        tot_count = result_today[5]
        print(response_data[period])
        if tot_count > 0:
            response_data[period]["tot_count"] = tot_count
            response_data[period]["yes_per"] = round(
                result_today[1] / tot_count * 100)
            response_data[period]["no_per"] = round(
                result_today[2] / tot_count * 100)
            response_data[period]["spoof_per"] = round(
                result_today[3] / tot_count * 100)
            response_data[period]["mrfault_per"] = round(
                result_today[4] / tot_count * 100)

    if result_week:
        period = result_week[0]
        tot_count = result_week[5]
        if tot_count > 0:
            response_data[period]["tot_count"] = tot_count
            response_data[period]["yes_per"] = round(
                result_week[1] / tot_count * 100)
            response_data[period]["no_per"] = round(
                result_week[2] / tot_count * 100)
            response_data[period]["spoof_per"] = round(
                result_week[3] / tot_count * 100)
            response_data[period]["mrfault_per"] = round(
                result_week[4] / tot_count * 100)

    for data in result_rows:
        period = data[0]
        tot_count = data[5]
        if tot_count > 0:
            response_data[period]["tot_count"] = tot_count
            response_data[period]["yes_per"] = round(data[1] / tot_count * 100)
            response_data[period]["no_per"] = round(data[2] / tot_count * 100)
            response_data[period]["spoof_per"] = round(
                data[3] / tot_count * 100)
            response_data[period]["mrfault_per"] = round(
                data[4] / tot_count * 100)

    return Response({
        "status": bool(response_data),
        "message": "Data Fetched Successfully" if response_data else "Data not found",
        "data": response_data,
    })

@api_view(['GET', 'POST'])
def getuserdata(request):
    user = request.data.get('user', None)
    month = request.data.get('month', None)
    year = request.data.get('year', None)
    print("useryear", year)
    now = datetime.now()
    cursor = connection.cursor()

    if user:
        user_condition = f"AND mtr_excep_img = 'vapp_{user}'"
        user_counts = f"COUNT(CASE WHEN mtr_excep_img = 'vapp_{user}' THEN 1 END) AS user_count"
    else:
        user_condition = ""
        user_counts = ', '.join(
            [f"COUNT(CASE WHEN mtr_excep_img = 'vapp_user{i}' THEN 1 END) AS user{i}" for i in range(1, 16)])
    if month:
        month_condition = f"extract(month from date_qc::date) = '{month}'"
    else:
        month_condition = f"extract(month from date_qc::date) = '{now.month}'"
    if year:
        year_condition = f"extract(year from date_qc::date) = '{year}'"
    else:
        year_condition = f"extract(year from date_qc::date) = '{now.year}'"
    query = f'''
    SELECT
        {user_counts},
        count(*) as tot_count
    FROM readingmaster
    WHERE 
       {month_condition} AND {year_condition}
        AND date_qc <> ''
        {user_condition}
    '''

    print("query", query)
    cursor.execute(query)
    person_objects = dictfetchall(cursor)

    if not person_objects:
        return Response({
            "status": False,
            "message": "Data not found",
            "data": person_objects
        })

    response_data = {
        "status": True,
        "message": "Data Fetched Successfully",
        "data": person_objects
    }

    return Response(response_data)


@api_view(['POST', 'GET'])
def downloadmrlist(request):
    print("request", request.query_params)
    now = datetime.now()
    month = request.query_params.get('month', now.month)
    year = request.query_params.get('year', now.year)
    image_type = request.query_params.get('image_type',  None)
    cursor = connection.cursor()

    if image_type == 'Not Found':
        ocr_reading = f"prsnt_ocr_rdng = 'Not Found'"
    elif image_type == 'Found':
        ocr_reading = f"prsnt_ocr_rdng != 'Not Found'"
    query = f'''
    select mr_id, count(*) as tot_count from readingmaster WHERE rdng_ocr_status = 'Failed'
and prsnt_rdng_ocr_excep != 'Spoofed Image' and manual_update_flag is null and is_object_meter!='NO and camera blocked'
AND EXTRACT(MONTH FROM reading_date_db)='{month}' AND EXTRACT(YEAR FROM reading_date_db)='{year}' AND {ocr_reading} 
group by mr_id order by tot_count desc
    '''
    print("query", query)
    cursor.execute(query)
    person_objects = dictfetchall(cursor)
    print("data", person_objects)

    if not person_objects:
        return Response({
            "status": False,
            "message": "Data not found",
            "data": person_objects
        })

    response_data = {
        "status": True,
        "message": "Data Fetched Successfully",
        "data": person_objects
    }

    return Response(response_data)


@api_view(['POST', 'GET'])
def downloaddivisionlist(request):
    print("request", request.query_params)
    now = datetime.now()
    month = request.query_params.get('month', now.month)
    year = request.query_params.get('year', now.year)
    discom = request.query_params.get('discom', 'NBPDCL')
    zone = request.query_params.get('zone', 'NORTH BIHAR RURAL')
    image_type = request.query_params.get('image_type',  None)
    cursor = connection.cursor()

    if image_type == 'Not Found':
        ocr_reading = f"prsnt_ocr_rdng = 'Not Found'"
    elif image_type == 'Found':
        ocr_reading = f"prsnt_ocr_rdng != 'Not Found'"
    query = f'''
    select ofc_division, count(*) as tot_count from readingmaster WHERE rdng_ocr_status = 'Failed'
and prsnt_rdng_ocr_excep != 'Spoofed Image' and manual_update_flag is null and is_object_meter!='NO and camera blocked'
AND EXTRACT(MONTH FROM reading_date_db)='{month}' AND EXTRACT(YEAR FROM reading_date_db)='{year}' AND ofc_discom = '{discom}' AND ofc_zone = '{zone}' AND {ocr_reading} 
group by ofc_division order by tot_count desc
    '''
    print("query", query)
    cursor.execute(query)
    person_objects = dictfetchall(cursor)
    print("data", person_objects)

    if not person_objects:
        return Response({
            "status": False,
            "message": "Data not found",
            "data": person_objects
        })

    response_data = {
        "status": True,
        "message": "Data Fetched Successfully",
        "data": person_objects
    }

    return Response(response_data)


@api_view(["POST"])
def originalimageApi(request):
    data = request.data.copy()
    mr_id = data.get("mr_id")
    mr_ids = []

    if mr_id in mr_ids:
        return Response({"status": True})
    else:
        return Response({"status": False})

@api_view(["GET", "POST"])
def meter_reading_summary_new(request):
    mrid = request.data.get("mrid", None)
    agency = request.data.get("agency", None)
    ofc_zone = request.data.get("ofc_zone", None)
    month = request.data.get("month", None)

    # to check current month and previous month
    current_month = date.today().month
    previous_month = (date.today() - timedelta(days=30)).month

    # decide which table to use based on the month
    if month:
        if month == f"{current_month:02d}" or month == f"{previous_month:02d}":
            table_name = "readingmaster"
        else:
            table_name = "prevmonthsdata"
    else:
        table_name = "readingmaster"
    clause = f" WHERE "
    clause += (
        f" EXTRACT(MONTH from reading_date_db)='{month.split('-')[0]}' "
        if (month)
        else f" EXTRACT(MONTH from reading_date_db)='{current_month}' "
    )
    clause += f" AND bl_agnc_name='{agency}' " if (agency) else ""
    clause += f" AND r.ofc_zone = '{ofc_zone}' " if (ofc_zone) else ""
    cursor = connection.cursor()
    query = f""" select  r.mr_id, count(r.mr_id),
    count(r.prsnt_mtr_status='Ok' OR NULL) as ok_readings,
    count(r.rdng_ocr_status='Passed' or NULL) as OCRwithoutException,
    count(r.rdng_ocr_status='Failed' OR NULL) as OCRwithException,
    count(r.prsnt_mtr_status='Meter Defective' OR NULL) as MeterDefective,
    count(r.prsnt_mtr_status='Door Locked' OR NULL) as DoorLocked
    from {table_name} r {clause} and r.mr_id!='' group by  r.mr_id
    """
    print("Query--------", query)
    location = []
    count = 0
    try:
        cursor.execute(query)
        res = cursor.fetchall()
        for row in res:
            mrid = row[0]
            total = row[1]
            ok_readings = row[2]
            ocr_passed = row[3]
            ocr_failed = row[4]
            md = row[5]
            dl = row[6]
            location.append(
                {
                    "mrid": mrid,
                    "totalReadings": total,
                    "OKreadings": ok_readings,
                    "OCRReadings": ocr_passed,
                    "OCRwithException": ocr_failed,
                    "MeterDefective": md,
                    "DoorLocked": dl,
                    "OKreadingspercent": math.floor(
                        ((ok_readings / total) if ok_readings else 0) * 100
                    ),
                    "OCRReadingspercent": math.floor(
                        ((ocr_passed / ok_readings) if ocr_passed else 0) * 100
                    ),
                    "OCRwithExceptionpercent": math.floor(
                        ((ocr_failed / ok_readings) if ocr_failed else 0) * 100
                    ),
                    "MeterDefectivepercent": math.floor(
                        ((md / total) if md else 0) * 100
                    ),
                    "DoorLockedpercent": math.floor(((dl / total) if dl else 0) * 100),
                }
            )

            count += 1

        return Response(location)
    except Exception as e:
        print("Exception----------", e)
        return Response([])

   
# --------------------------------------------------------------------------------------------------------------------------------
# MATERIALIZED VIEWS APIS
# from datetime import date
from django.db import connection
from rest_framework.decorators import api_view
from rest_framework.response import Response

@api_view(["GET"])
def refreshAPI(request):
    cursor = connection.cursor()
    updateddate = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    query = """ 
      REFRESH MATERIALIZED VIEW exception_material_view1;
      REFRESH MATERIALIZED VIEW currentdate_materialized_view;
      REFRESH MATERIALIZED VIEW currentmonth_materialized_view;
      REFRESH MATERIALIZED VIEW previousmonth_materialized_view;

      """
    cursor.execute(query)
    return Response({"msg": "Updated Succesfully", "updateddate": updateddate})

@api_view(["GET"])
def get_officedata(request):
    with open("D:/123.json") as f:
        data = json.load(f)
        for i in data:
            data = Office.objects.create(
                id=i["id"],
                discom=i["discom"],
                zone=i["zone"],
                circlename=i["circlename"],
                divisionname=i["divisionname"],
                divisioncode=i["divisioncode"],
                subdivision=i["subdivision"],
                subdivisioncode=i["subdivisioncode"],
                sectionname=i["sectionname"],
                sectioncode=i["sectioncode"],
                agency=i["agency"],
                agencycode=i["agencycode"],
            )
            data.save()
    return Response("data added Successfully")


@api_view(["GET"])
def getofficedata(request):
    cursor = connection.cursor()
    query = f"""
        select * from office   
"""
    print("QUERY-->", query)
    cursor.execute(query)
    result = dictfetchall(cursor)
    return Response(result)
